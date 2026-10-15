# -*- coding: utf-8 -*-
"""
Marmara Bölge Yönetim Platformu
Güzel Enerji · Marmara Bölge Müdürlüğü — proje, gelişme, ziyaret, bayi ve pazar payı takibi.
"""
import io, json, os, secrets, shutil, sqlite3, tempfile, datetime as dt
from collections import defaultdict
from flask import (Flask, render_template, request, redirect, url_for, flash, session, abort,
                   jsonify, send_file, g)
from werkzeug.security import check_password_hash, generate_password_hash
import db as vt
import hesap as H
import guvenlik as G

KOK = os.path.dirname(os.path.abspath(__file__))


def uygulama_olustur(test_yapilandirma=None):
    app = Flask(__name__, instance_path=os.path.join(KOK, "instance"))
    os.makedirs(app.instance_path, exist_ok=True)
    anahtar_yolu = os.path.join(app.instance_path, "gizli_anahtar")
    if not os.path.exists(anahtar_yolu):
        with open(anahtar_yolu, "w") as f:
            f.write(secrets.token_hex(32))
    app.config.update(
        SECRET_KEY=os.environ.get("GIZLI_ANAHTAR") or open(anahtar_yolu).read().strip(),
        VERITABANI=os.environ.get("VERITABANI") or os.path.join(app.instance_path, "platform.db"),
        SEED=os.path.join(KOK, "data", "seed_data.json"),
        SIFRE_DOSYASI=os.path.join(app.instance_path, "ilk_sifreler.txt"),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.environ.get("GUVENLI_CEREZ") == "1",
        PERMANENT_SESSION_LIFETIME=dt.timedelta(hours=10),
        MAX_CONTENT_LENGTH=2 * 1024 * 1024,
        ORNEK_VERI=os.environ.get("ORNEK_VERI", "1") == "1",
        HAZIR_KULLANICILAR=os.environ.get("HAZIR_KULLANICILAR", "0") == "1",
        SABIT_SIFRE=None,
        # Test modu: giriş ekranı atlanır, oturum bölge müdürü olarak açılır.
        # baslat.bat / baslat.sh bunu 1 yapar; sunucuya/buluta kurarken TEST_MODU=0 olmalıdır.
        TEST_MODU=os.environ.get("TEST_MODU", "0") == "1",
    )
    if test_yapilandirma:
        app.config.update(test_yapilandirma)
    sifreler = vt.kur(app.config["VERITABANI"], app.config["SEED"], app.config["SIFRE_DOSYASI"],
                      app.config["ORNEK_VERI"], app.config["SABIT_SIFRE"], app.config["HAZIR_KULLANICILAR"])
    if sifreler and not app.config.get("TESTING"):
        print("\n" + "=" * 64)
        print("  İLK KURULUM TAMAMLANDI — bölge müdürü ilk giriş şifresi:")
        for k, v in sifreler.items():
            print(f"    {k:10s} {v}")
        print(f"  (Aynı liste: {app.config['SIFRE_DOSYASI']})")
        print("=" * 64 + "\n")
    if app.config["TEST_MODU"] and not app.config.get("TESTING"):
        print("  ! TEST MODU AÇIK — giriş ekranı atlanıyor, herkes bölge müdürü olarak giriyor.")
        print("    Kapatmak için baslat dosyasındaki TEST_MODU satırını 0 yapın.\n")
    app.teardown_appcontext(vt.kapat)
    app.jinja_env.filters["ad"] = H.guzel_ad
    rotalari_kaydet(app)
    return app


ROL_ADI = {"admin": "Bölge Müdürü", "saha": "Saha Müdürü", "ziyaretci": "Ziyaretçi"}


# ── biçimlendirme yardımcıları ──────────────────────────────────────────────
def sayi(v, ond=0):
    if v is None or v == "":
        return "–"
    try:
        s = f"{float(v):,.{ond}f}"
    except (TypeError, ValueError):
        return str(v)
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def yuzde(v, ond=1):
    if v is None or v == "":
        return "–"
    return "%" + sayi(float(v) * 100, ond)


def bp(v):
    if v is None:
        return "–"
    v = round(v)
    return f"{'+' if v > 0 else ''}{v} bp"


ROZET = {
    "Geciken": "kirmizi", "Gecikmiş Plan": "kirmizi", "Kritik": "kirmizi", "Pay kaybetti": "kirmizi",
    "Engel": "kirmizi", "Bugün": "turuncu", "Bu Hafta": "turuncu", "Yüksek": "turuncu",
    "Bu Ay": "sari", "Orta": "sari", "Beklemede": "sari", "Gelecek Hafta": "mavi", "Karar": "mavi",
    "Devam Ediyor": "mavi", "Planlandı": "mavi", "Tamamlandı": "yesil", "Gerçekleşti": "yesil",
    "Pay kazandı": "yesil", "Temiz": "yesil", "Kapandı": "yesil", "Projeye dönüştü": "mavi",
    "İptal": "gri", "Ertelendi": "gri", "Bilgi": "gri", "Düşük": "gri", "İleride": "gri",
    "Yatay": "gri", "Tarihsiz": "gri", "Yapılacak": "gri", "Açık": "turuncu",
    "Yeni": "turuncu", "İnceleniyor": "sari", "Reddedildi": "gri", "Kapalı": "gri", "Arşiv": "gri",
    "Onaylandı": "yesil", "Koşullu onay": "sari", "Onaylanmadı / beklet": "kirmizi", "Not / talimat": "mavi",
    "Yanıt bekliyor": "turuncu", "İzlenecek": "sari",
}


def gelisme_ozet(d):
    ozet = {}
    for r in d.execute("SELECT proje_id, COUNT(*) n, MAX(tarih) son FROM gelismeler GROUP BY proje_id"):
        ozet[r["proje_id"]] = {"n": r["n"], "son": r["son"]}
    for r in d.execute("SELECT g.proje_id, g.metin FROM gelismeler g JOIN (SELECT proje_id, MAX(tarih) t,"
                       " MAX(id) mid FROM gelismeler GROUP BY proje_id) s ON g.proje_id=s.proje_id AND"
                       " g.tarih=s.t ORDER BY g.id"):
        if r["proje_id"] in ozet:
            ozet[r["proje_id"]]["metin"] = r["metin"]
    return ozet


def tum_projeler(d, arsiv="aktif"):
    """arsiv: 'aktif' (varsayılan — arşivdekiler hariç), 'arsiv' (yalnız arşiv), 'tumu'."""
    oz = gelisme_ozet(d)
    t = H.bugun()
    where = {"aktif": " WHERE COALESCE(arsiv,0)=0", "arsiv": " WHERE arsiv=1"}.get(arsiv, "")
    return [H.proje_zenginlestir(p, t, oz) for p in d.execute("SELECT * FROM projeler" + where)]


def tum_talepler(d, where="", params=()):
    t = H.bugun()
    rows = d.execute("SELECT t.*, b.unvan, b.il, b.ilce, (SELECT COUNT(*) FROM talep_notlari n WHERE n.talep_id=t.id) not_sayisi"
                     " FROM talepler t LEFT JOIN bayiler b ON b.id=t.bayi_id " + where +
                     " ORDER BY CASE t.durum WHEN 'Yeni' THEN 0 WHEN 'İnceleniyor' THEN 1 WHEN 'Devam Ediyor' THEN 2"
                     " ELSE 3 END, COALESCE(t.termin,'9999'), t.id DESC", params)
    return [H.talep_zenginlestir(r, t) for r in rows]


def kredi_hesapla(d, sm=None):
    """Borç raporu + onay yazışmalarından kredi-risk özetini çıkarır."""
    where, prm = ("WHERE sm_adi=?", (sm,)) if sm else ("", ())
    B = [H.borc_zenginlestir(r) for r in d.execute("SELECT * FROM borc " + where, prm)]
    borclu = [b for b in B if (b["vadesi_gecmis"] or 0) > 0]
    top = lambda L, k: sum((x[k] or 0) for x in L)
    rapor = d.execute("SELECT rapor_tarihi, onceki_tarih FROM borc LIMIT 1").fetchone()
    O = [dict(r) for r in d.execute("SELECT * FROM risk_onaylari " + where + " ORDER BY zaman DESC", prm)]
    onayli = [o for o in O if o["onay_durumu"] in ("Onaylandı", "Koşullu onay")]
    sure = sorted(o["sure_dk"] for o in O if o["sure_dk"] is not None)
    sm_tablo = []
    for s_ in ([sm] if sm else H.SM_ADLARI):
        L = [b for b in B if b["sm_adi"] == s_]
        BL = [b for b in L if (b["vadesi_gecmis"] or 0) > 0]
        OL = [o for o in O if o["sm_adi"] == s_]
        vg = top(BL, "vadesi_gecmis")
        sm_tablo.append({
            "sm": s_, "kod": H.SM_KOD.get(s_, ""), "bayi": len(L), "borclu": len(BL), "vadesi_gecmis": vg,
            "fark": top(L, "fark"), "eskiyen": top(BL, "eskiyen"),
            "eskiyen_pay": (top(BL, "eskiyen") / vg) if vg else 0,
            "acik_hesap": top(L, "acik_hesap"), "cek_senet": top(L, "cek_senet"),
            "toplam_risk": top(L, "toplam_risk"), "kritik": sum(1 for b in BL if b["seviye"] == "Kritik"),
            "onay": len(OL), "onay_son30": sum(1 for o in OL if o["zaman"] >= H.bugun().replace(day=1).isoformat()),
            **{k: top(BL, k) for k, _ in H.YAS_KOLONLARI},
        })
    ozet = {
        "vadesi_gecmis": top(B, "vadesi_gecmis"), "fark": top(B, "fark"), "borclu": len(borclu),
        "bayi": len(B), "eskiyen": top(B, "eskiyen"), "acik_hesap": top(B, "acik_hesap"),
        "cek_senet": top(B, "cek_senet"), "toplam_risk": top(B, "toplam_risk"),
        "kritik": sum(1 for b in borclu if b["seviye"] == "Kritik"),
        "rapor_tarihi": rapor["rapor_tarihi"] if rapor else None,
        "onceki_tarih": rapor["onceki_tarih"] if rapor else None,
        "onay": len(O), "onay_oran": (len(onayli) / len(O)) if O else 0,
        "medyan_dk": sure[len(sure) // 2] if sure else None,
        **{k: top(B, k) for k, _ in H.YAS_KOLONLARI},
    }
    borclu.sort(key=lambda b: -b["skor"])
    return {"ozet": ozet, "sm": sm_tablo, "borclu": borclu, "tumu": B, "onaylar": O}


def bayi_riskleri(d):
    risk = {}
    for r in d.execute("SELECT bayi_id, onem FROM aksakliklar WHERE durum='Açık' AND bayi_id IS NOT NULL"):
        mevcut = risk.get(r["bayi_id"])
        if mevcut is None or H.ONEM_SIRA.get(r["onem"], 9) < H.ONEM_SIRA.get(mevcut, 9):
            risk[r["bayi_id"]] = r["onem"]
    return risk


def tum_ziyaretler(d, where="", params=()):
    t = H.bugun()
    rows = d.execute("SELECT z.*, b.unvan, b.il, b.ilce, b.segment, b.ist_kodu, b.ay_hacim FROM ziyaretler z"
                     " LEFT JOIN bayiler b ON b.id=z.bayi_id " + where + " ORDER BY z.tarih, z.id", params)
    out = []
    for r in rows:
        z = dict(r)
        z["zaman"] = H.ziyaret_zaman_durumu(z, t)
        out.append(z)
    return out


def pazar_payi_hesapla(d):
    il = {r["il"]: dict(r) for r in d.execute("SELECT * FROM pazar_il")}
    agirlik = defaultdict(dict)
    il_toplam = defaultdict(int)
    for r in d.execute("SELECT * FROM sm_il_istasyon"):
        agirlik[r["sm_adi"]][r["il"]] = r["adet"]
        il_toplam[r["il"]] += r["adet"]
    sehirler = []
    for c in H.MARMARA_ILLERI:
        o = il.get(c)
        if not o:
            continue
        row = H.pp_satiri(o["ge25"], o["ge26"], o["s25"], o["s26"])
        row["il"] = c
        row["istasyon"] = il_toplam.get(c, 0)
        row["sahipler"] = [(s, agirlik[s][c]) for s in H.SM_ADLARI if agirlik[s].get(c)]
        sehirler.append(row)
    sehirler.sort(key=lambda x: -x["bp"])
    smler = []
    for s in H.SM_ADLARI:
        a = {k: 0.0 for k in ("ge25", "ge26", "s25", "s26")}
        for c, n in agirlik[s].items():
            o = il.get(c)
            if not o or not il_toplam[c]:
                continue
            w = n / il_toplam[c]
            for k in a:
                a[k] += w * o[k]
        row = H.pp_satiri(a["ge25"], a["ge26"], a["s25"], a["s26"])
        row["sm"] = s
        row["istasyon"] = sum(agirlik[s].values())
        row["iller"] = [(c, n) for c, n in agirlik[s].items()]
        smler.append(row)
    top = {k: sum(il[c][k] for c in H.MARMARA_ILLERI if c in il) for k in ("ge25", "ge26", "s25", "s26")}
    bolge = H.pp_satiri(**top)
    bolge["istasyon"] = sum(il_toplam.values())
    tr = il.get("Türkiye")
    turkiye = H.pp_satiri(tr["ge25"], tr["ge26"], tr["s25"], tr["s26"]) if tr else None
    if turkiye:
        turkiye["istasyon"] = tr["ist_epdk"]
    aylik = []
    for i in range(7):
        v = {k: 0.0 for k in ("ge25", "ge26", "s25", "s26")}
        for c in H.MARMARA_ILLERI:
            o = il.get(c)
            if not o:
                continue
            v["ge25"] += json.loads(o["ge_ay25"])[i]; v["ge26"] += json.loads(o["ge_ay26"])[i]
            v["s25"] += json.loads(o["s_ay25"])[i]; v["s26"] += json.loads(o["s_ay26"])[i]
        row = H.pp_satiri(**v)
        row["ay"] = H.AYLAR[i]
        aylik.append(row)
    return {"bolge": bolge, "turkiye": turkiye, "sm": smler, "sehir": sehirler, "aylik": aylik}


def mintika_hesapla(d):
    bayiler = [dict(r) for r in d.execute("SELECT * FROM bayiler")]
    aks = [dict(r) for r in d.execute("SELECT * FROM aksakliklar")]
    toplam_hacim = sum(b["ay_hacim"] or 0 for b in bayiler) or 1
    sonuc = []
    for s in H.SM_ADLARI + ["BÖLGE"]:
        B = bayiler if s == "BÖLGE" else [b for b in bayiler if b["sm_adi"] == s]
        A = [a for a in aks if (s == "BÖLGE" or a["sm_adi"] == s) and a["durum"] == "Açık"]
        n = len(B) or 1
        hacim = sum(b["ay_hacim"] or 0 for b in B)
        ytd26 = sum(b["ytd26_hacim"] or 0 for b in B)
        ytd25 = sum(b["ytd25_hacim"] or 0 for b in B)
        kar = sum(b["ytd26_kar"] or 0 for b in B)
        enb = max((b["ay_hacim"] or 0 for b in B), default=0)
        zarar = [b for b in B if (b["ytd26_kar"] or 0) < 0]
        sonuc.append({
            "sm": s, "istasyon": len(B), "il": len({b["il"] for b in B}), "ilce": len({b["ilce"] for b in B}),
            "total": sum(1 for b in B if b["marka"] == "TotalEnergies"),
            "moil": sum(1 for b in B if b["marka"] == "MOİL"),
            "dodo": sum(1 for b in B if b["tip"] == "DODO"),
            "otogazli": sum(1 for b in B if (b["otogaz"] or 0) > 0),
            "hacim": hacim, "gunluk": hacim / 31, "ist_basi": hacim / n,
            "en_buyuk": enb, "en_buyuk_pay": enb / hacim if hacim else 0,
            "a_seg": sum(1 for b in B if b["segment"] == "A"), "bolge_payi": hacim / toplam_hacim,
            "ytd26": ytd26, "kar": kar, "kar_m3": kar / ytd26 if ytd26 else 0,
            "zarar_n": len(zarar), "zarar_tl": sum(b["ytd26_kar"] for b in zarar),
            "yoy": ytd26 / ytd25 - 1 if ytd25 else 0,
            "hic_satis": sum(1 for b in B if b["sifir_gun"] == 31),
            "sifir_gunlu": sum(1 for b in B if (b["sifir_gun"] or 0) > 0),
            "dusen": sum(1 for b in B if b["ytd25_hacim"] and b["ytd26_hacim"] / b["ytd25_hacim"] - 1 < -0.15),
            "kritik": sum(1 for a in A if a["onem"] == "Kritik"),
            "yuksek": sum(1 for a in A if a["onem"] == "Yüksek"),
            "orta": sum(1 for a in A if a["onem"] == "Orta"),
            "bulgu": len(A), "ist_bulgu": len(A) / n,
        })
    return sonuc


def rotalari_kaydet(app):

    @app.before_request
    def _once():
        G.csrf_kontrol()
        G.ziyaretci_kontrol()

    @app.after_request
    def _basliklar(resp):
        resp.headers["X-Frame-Options"] = "DENY"
        resp.headers["X-Content-Type-Options"] = "nosniff"
        resp.headers["Referrer-Policy"] = "same-origin"
        resp.headers["Content-Security-Policy"] = ("default-src 'self'; img-src 'self' data:; "
                                                   "style-src 'self' 'unsafe-inline'; script-src 'self'")
        return resp

    @app.context_processor
    def _baglam():
        u = G.aktif_kullanici()
        bekleyen = 0
        yeni_is = 0
        if u and u["rol"] == "admin":
            bekleyen = vt.baglan().execute("SELECT COUNT(*) FROM kullanicilar WHERE onay_durumu='onay_bekliyor'").fetchone()[0]
            yeni_is = vt.baglan().execute("SELECT COUNT(*) FROM talepler WHERE durum='Yeni'").fetchone()[0]
        test_kullanicilar = []
        if app.config["TEST_MODU"]:
            test_kullanicilar = [dict(r) for r in vt.baglan().execute(
                "SELECT id, ad_soyad, rol, sm_kodu FROM kullanicilar WHERE aktif=1 ORDER BY rol, sm_kodu")]
        return dict(ku=u, H=H, test_modu=app.config["TEST_MODU"], test_kullanicilar=test_kullanicilar, csrf=G.csrf_token, sayi=sayi, yuzde=yuzde, bp=bp, trt=H.tr_tarih, bekleyen_kayit=bekleyen, yeni_is=yeni_is,
                    rozet=lambda v: ROZET.get(v, "gri"), duzenleyebilir=G.duzenleyebilir,
                    salt_okunur=bool(u and u["rol"] == "ziyaretci"), ROL_ADI=ROL_ADI,
                    bugun=H.bugun())

    @app.errorhandler(403)
    def _e403(e):
        u = G.aktif_kullanici()
        if u and u["rol"] == "ziyaretci":
            return render_template("hata.html", kod=403, mesaj="Ziyaretçi hesabı yalnızca görüntüleme yetkisine sahiptir; "
                                   "kayıt ekleyemez veya değiştiremez."), 403
        return render_template("hata.html", kod=403, mesaj="Bu işlem için yetkiniz yok. Saha müdürleri yalnızca "
                               "kendi mıntıkalarına ait kayıtları değiştirebilir."), 403

    @app.errorhandler(404)
    def _e404(e):
        return render_template("hata.html", kod=404, mesaj="Aradığınız kayıt bulunamadı."), 404

    @app.errorhandler(400)
    def _e400(e):
        return render_template("hata.html", kod=400, mesaj=getattr(e, "description", "Geçersiz istek.")), 400

    # ── oturum ──────────────────────────────────────────────────────────────
    @app.route("/giris", methods=["GET", "POST"])
    def giris():
        if app.config["TEST_MODU"]:
            return redirect(url_for("pano"))
        if request.method == "POST":
            ka = request.form.get("kullanici_adi", "").strip().lower()
            sf = request.form.get("sifre", "")
            d = vt.baglan()
            r = d.execute("SELECT * FROM kullanicilar WHERE kullanici_adi=?", (ka,)).fetchone()
            u = dict(r) if r else None
            if u and G.kilitli_mi(u):
                flash("Çok sayıda hatalı deneme yapıldı. Hesap birkaç dakikalığına kilitlendi.", "hata")
            elif u and check_password_hash(u["sifre_hash"], sf) and u.get("onay_durumu") == "onay_bekliyor":
                flash("Kaydınız alındı ancak henüz bölge müdürü tarafından onaylanmadı. Onaylandıktan sonra giriş yapabilirsiniz.", "uyari")
            elif u and check_password_hash(u["sifre_hash"], sf) and u.get("onay_durumu") == "reddedildi":
                flash("Kayıt başvurunuz onaylanmadı. Bölge müdürüyle görüşün.", "hata")
            elif u and check_password_hash(u["sifre_hash"], sf) and not u["aktif"]:
                flash("Bu hesap pasif durumda. Bölge müdürüyle görüşün.", "hata")
            elif u and check_password_hash(u["sifre_hash"], sf):
                session.clear()
                session.permanent = True
                session["uid"] = u["id"]
                d.execute("UPDATE kullanicilar SET hatali_giris=0, kilit_bitis=NULL, son_giris=? WHERE id=?",
                          (vt.simdi(), u["id"]))
                vt.logla(u["kullanici_adi"], "giriş", "oturum")
                d.commit()
                hedef = request.args.get("sonra") or url_for("pano")
                if not hedef.startswith("/") or hedef.startswith("//"):
                    hedef = url_for("pano")
                return redirect(hedef)
            else:
                if u:
                    kilit = G.hatali_giris_kaydet(u)
                    if kilit:
                        flash("Çok sayıda hatalı deneme. Hesap 10 dakika kilitlendi.", "hata")
                        return render_template("giris.html")
                flash("Kullanıcı adı veya şifre hatalı.", "hata")
        return render_template("giris.html")

    KAYIT_DENEME = defaultdict(list)

    def _mintikalar():
        iller = defaultdict(list)
        for r in vt.baglan().execute("SELECT sm_adi, il, adet FROM sm_il_istasyon ORDER BY adet DESC"):
            iller[r["sm_adi"]].append(r["il"])
        return [{"kod": k, "sm": s, "iller": iller.get(s, [])} for s, k in H.SAHA_MUDURLERI]

    def _bm_hesabi():
        """Kurulumda açılan, henüz sahiplenilmemiş bölge müdürü hesabı (ilk şifresi değiştirilmemiş)."""
        return vt.baglan().execute("SELECT * FROM kullanicilar WHERE rol='admin' AND sm_kodu='NW1' AND sifre_degismeli=1"
                                   " AND aktif=1 ORDER BY id LIMIT 1").fetchone()

    def _kayit_sayfa(fo):
        return render_template("kayit.html", fo=fo, mintikalar=_mintikalar(), bm_bos=_bm_hesabi() is not None)

    @app.route("/kayit", methods=["GET", "POST"])
    def kayit():
        if G.aktif_kullanici():
            return redirect(url_for("pano"))
        fo = {"ad_soyad": "", "mintika": request.args.get("tur", ""), "kullanici_adi": ""}
        if request.method == "POST":
            ip = request.remote_addr or "-"
            simdi_ = dt.datetime.now()
            KAYIT_DENEME[ip] = [t for t in KAYIT_DENEME[ip] if (simdi_ - t).seconds < 3600]
            fo = {k: (request.form.get(k) or "").strip() for k in ("ad_soyad", "mintika", "kullanici_adi")}
            fo["kullanici_adi"] = fo["kullanici_adi"].lower()
            s1, s2 = request.form.get("sifre", ""), request.form.get("sifre2", "")
            kodlar = {k: s for s, k in H.SAHA_MUDURLERI}
            bm = fo["mintika"] == "NW1"
            zy = fo["mintika"] == "ZIYARETCI"
            bm_hesap = _bm_hesabi() if bm else None
            hatalar = []
            if len(KAYIT_DENEME[ip]) >= 10:
                hatalar.append("Çok fazla kayıt denemesi yapıldı. Lütfen daha sonra tekrar deneyin.")
            if fo["mintika"] not in kodlar and not bm and not zy:
                hatalar.append("Kullanıcı türünüzü / mıntıkanızı seçin.")
            if bm:
                fo["ad_soyad"] = "Erhan Başaçar"
                if not bm_hesap:
                    hatalar.append("Bölge müdürü hesabı zaten oluşturulmuş. Giriş sayfasından giriş yapın; "
                                   "şifrenizi unuttuysanız README'deki 'Şifre sıfırlama' adımını izleyin.")
                elif not check_password_hash(bm_hesap["sifre_hash"], request.form.get("kurulum_kodu", "").strip()):
                    hatalar.append("Kurulum kodu hatalı. Kod, platformu ilk başlattığınızda siyah pencerede ve "
                                   "instance/ilk_sifreler.txt dosyasında yazan şifredir.")
            elif len(fo["ad_soyad"]) < 3:
                hatalar.append("Adınızı ve soyadınızı yazın.")
            import re as _re
            if not _re.fullmatch(r"[a-z0-9._]{3,20}", fo["kullanici_adi"]):
                hatalar.append("Kullanıcı adı 3–20 karakter olmalı; yalnızca küçük harf, rakam, nokta ve alt çizgi kullanılabilir.")
            elif vt.baglan().execute("SELECT 1 FROM kullanicilar WHERE kullanici_adi=? AND id!=?",
                                     (fo["kullanici_adi"], bm_hesap["id"] if bm_hesap else -1)).fetchone():
                hatalar.append("Bu kullanıcı adı alınmış, başka bir tane seçin.")
            if s1 != s2:
                hatalar.append("Şifreler birbiriyle aynı değil.")
            else:
                h = G.sifre_uygun(s1)
                if h: hatalar.append(h)
            KAYIT_DENEME[ip].append(simdi_)
            if hatalar:
                for h in hatalar: flash(h, "hata")
                return _kayit_sayfa(fo)
            d = vt.baglan()
            if bm:
                # Kurulum hesabını sahiplen: kullanıcı adı ve şifre Erhan Bey'in seçtiği olur, onay gerekmez
                d.execute("UPDATE kullanicilar SET kullanici_adi=?, ad_soyad=?, sifre_hash=?, sifre_degismeli=0,"
                          " hatali_giris=0, kilit_bitis=NULL, onay_durumu='onaylı' WHERE id=?",
                          (fo["kullanici_adi"], "Erhan Başaçar", generate_password_hash(s1), bm_hesap["id"]))
                vt.logla(fo["kullanici_adi"], "bölge müdürü kaydı", "kullanıcı", fo["kullanici_adi"],
                         f"kurulum hesabı ({bm_hesap['kullanici_adi']}) sahiplenildi")
                d.commit()
                flash("Bölge müdürü hesabınız hazır. Belirlediğiniz kullanıcı adı ve şifre ile giriş yapın.", "basari")
                return redirect(url_for("giris"))
            if zy:
                rol, sm, kod = "ziyaretci", None, None
            else:
                rol, sm, kod = "saha", kodlar[fo["mintika"]], fo["mintika"]
            d.execute("INSERT INTO kullanicilar (kullanici_adi, ad_soyad, rol, sm_adi, sm_kodu, sifre_hash, sifre_degismeli,"
                      " aktif, onay_durumu, talep_mintika, olusturma) VALUES (?,?,?,?,?,?,0,0,'onay_bekliyor',?,?)",
                      (fo["kullanici_adi"], fo["ad_soyad"], rol, sm, kod, generate_password_hash(s1),
                       fo["mintika"], vt.simdi()))
            vt.logla(fo["kullanici_adi"], "kayıt başvurusu", "kullanıcı", fo["kullanici_adi"],
                     f"{'Ziyaretçi' if zy else fo['mintika']} · {fo['ad_soyad']}")
            d.commit()
            return render_template("kayit_bitti.html", fo=fo, sm=sm, ziyaretci=zy)
        return _kayit_sayfa(fo)

    @app.route("/test-kullanici/<int:kid>")
    def test_kullanici(kid):
        """Test modunda rolleri denemek için kullanıcı değiştirir."""
        if not app.config["TEST_MODU"]:
            abort(404)
        k = vt.baglan().execute("SELECT * FROM kullanicilar WHERE id=? AND aktif=1", (kid,)).fetchone()
        if not k:
            abort(404)
        session["uid"] = kid
        session["_test"] = 1
        flash(f"Test kullanıcısı: {k['ad_soyad']} ({ROL_ADI.get(k['rol'], k['rol'])})", "bilgi")
        return redirect(url_for("pano"))

    @app.route("/cikis")
    def cikis():
        u = G.aktif_kullanici()
        if u:
            vt.logla(u["kullanici_adi"], "çıkış", "oturum"); vt.baglan().commit()
        session.clear()
        flash("Oturum kapatıldı.", "bilgi")
        return redirect(url_for("giris"))

    @app.route("/sifre", methods=["GET", "POST"])
    @G.giris_gerekli
    def sifre_degistir():
        u = G.aktif_kullanici()
        if request.method == "POST":
            eski, yeni, yeni2 = (request.form.get(k, "") for k in ("eski", "yeni", "yeni2"))
            hata = None
            if not check_password_hash(u["sifre_hash"], eski):
                hata = "Mevcut şifre hatalı."
            elif yeni != yeni2:
                hata = "Yeni şifreler birbiriyle aynı değil."
            elif yeni == eski:
                hata = "Yeni şifre eskisiyle aynı olamaz."
            else:
                hata = G.sifre_uygun(yeni)
            if hata:
                flash(hata, "hata")
            else:
                d = vt.baglan()
                d.execute("UPDATE kullanicilar SET sifre_hash=?, sifre_degismeli=0 WHERE id=?",
                          (generate_password_hash(yeni), u["id"]))
                vt.logla(u["kullanici_adi"], "şifre değiştirme", "kullanıcı", u["kullanici_adi"])
                d.commit()
                flash("Şifreniz güncellendi.", "basari")
                return redirect(url_for("pano"))
        return render_template("sifre.html")

    # ── pano ────────────────────────────────────────────────────────────────
    @app.route("/")
    @G.giris_gerekli
    def pano():
        d = vt.baglan()
        u = G.aktif_kullanici()
        P = tum_projeler(d)
        Z = tum_ziyaretler(d)
        pp = pazar_payi_hesapla(d)
        t = H.bugun()
        acik = [p for p in P if p["durum"] not in ("Tamamlandı", "İptal")]
        yedi_gun = (t - dt.timedelta(days=7)).isoformat()
        gel7 = defaultdict(int)
        for r in d.execute("SELECT p.sm_adi, COUNT(*) n FROM gelismeler g JOIN projeler p ON p.id=g.proje_id"
                           " WHERE g.tarih>=? AND COALESCE(p.arsiv,0)=0 GROUP BY p.sm_adi", (yedi_gun,)):
            gel7[r["sm_adi"]] = r["n"]
        pp_sm = {r["sm"]: r for r in pp["sm"]}
        sm_tablo = []
        for s in H.SM_ADLARI:
            S = [p for p in P if p["sm_adi"] == s]
            A = [p for p in S if p["durum"] not in ("Tamamlandı", "İptal")]
            ZS = [z for z in Z if z["sm_adi"] == s]
            sm_tablo.append({
                "sm": s, "kod": H.SM_KOD[s], "toplam": len(S),
                "tamam": sum(1 for p in S if p["durum"] == "Tamamlandı"),
                "devam": sum(1 for p in S if p["durum"] == "Devam Ediyor"),
                "acik": len(A), "geciken": sum(1 for p in S if p["zaman"] == "Geciken"),
                "tamam_oran": (sum(1 for p in S if p["durum"] == "Tamamlandı") / len(S)) if S else 0,
                "ilerleme": (sum(p["ilerleme"] or 0 for p in S) / len(S)) if S else 0,
                "bu_hafta": sum(1 for p in S if p["zaman"] in ("Bu Hafta", "Bugün")),
                "kritik": sum(1 for p in A if p["oncelik"] in ("Kritik", "Yüksek")),
                "gel7": gel7[s], "sessiz": sum(1 for p in A if (p["temassiz"] or 0) >= 14),
                "z_hafta": sum(1 for z in ZS if z["zaman"] == "Bu Hafta"),
                "z_gecikmis": sum(1 for z in ZS if z["zaman"] == "Gecikmiş Plan"),
                "pp26": pp_sm[s]["pp26"], "pp_bp": pp_sm[s]["bp"],
            })
        kpi = {
            "acik": len(acik), "geciken": sum(1 for p in P if p["zaman"] == "Geciken"),
            "bu_hafta": sum(1 for p in P if p["zaman"] in ("Bu Hafta", "Bugün")),
            "tamam_oran": (sum(1 for p in P if p["durum"] == "Tamamlandı") / len(P)) if P else 0,
            "z_hafta": sum(1 for z in Z if z["zaman"] == "Bu Hafta"),
            "z_gecikmis": sum(1 for z in Z if z["zaman"] == "Gecikmiş Plan"),
            "pp26": pp["bolge"]["pp26"], "pp_bp": pp["bolge"]["bp"],
            "kritik_bulgu": d.execute("SELECT COUNT(*) FROM aksakliklar WHERE durum='Açık' AND onem='Kritik'")
            .fetchone()[0],
        }
        oncelikli = sorted([p for p in acik], key=lambda p: -p["skor"])[:15]
        isi = {s: {dp: 0 for dp in H.DEPARTMANLAR} for s in H.SM_ADLARI}
        for p in acik:
            if p["sm_adi"] in isi and p["departman"] in isi[p["sm_adi"]]:
                isi[p["sm_adi"]][p["departman"]] += (3 if p["zaman"] == "Geciken" else 0) + \
                    (2 if p["oncelik"] == "Kritik" else 0) + (1 if p["zaman"] == "Bu Hafta" else 0)
        hb = H.hafta_basi(t)
        hafta = []
        for i in range(7):
            gun = hb + dt.timedelta(days=i)
            liste = [z for z in Z if z["tarih"] == gun.isoformat() and z["durum"] != "İptal"]
            if i < 5 or liste:
                hafta.append({"gun": gun, "ad": H.GUNLER[i], "ziyaretler": liste, "bugun": gun == t})
        gecikmis_z = [z for z in Z if z["zaman"] == "Gecikmiş Plan"]
        benim = None
        if u["rol"] == "saha":
            benim = {
                "geciken": [p for p in acik if p["sm_adi"] == u["sm_adi"] and p["zaman"] == "Geciken"],
                "hafta": [p for p in acik if p["sm_adi"] == u["sm_adi"] and p["zaman"] in ("Bu Hafta", "Bugün")],
                "sessiz": [p for p in acik if p["sm_adi"] == u["sm_adi"] and (p["temassiz"] or 0) >= 14],
                "ziyaret": [z for z in Z if z["sm_adi"] == u["sm_adi"] and z["zaman"] in ("Bu Hafta", "Gecikmiş Plan")],
            }
        kredi = kredi_hesapla(d)
        kredi_sm = {x["sm"]: x for x in kredi["sm"]}
        for r in sm_tablo:
            k = kredi_sm.get(r["sm"], {})
            r["vadesi_gecmis"] = k.get("vadesi_gecmis", 0)
            r["eskiyen"] = k.get("eskiyen", 0)
            r["borclu"] = k.get("borclu", 0)
        kpi["vadesi_gecmis"] = kredi["ozet"]["vadesi_gecmis"]
        kpi["borc_fark"] = kredi["ozet"]["fark"]
        kpi["eskiyen"] = kredi["ozet"]["eskiyen"]
        kpi["borclu"] = kredi["ozet"]["borclu"]
        talep = None
        if u["rol"] == "admin":
            T = tum_talepler(d, "WHERE t.durum NOT IN ('Tamamlandı','Reddedildi')")
            talep = {"liste": T[:6], "acik": len(T), "yeni": sum(1 for x in T if x["durum"] == "Yeni"),
                     "geciken": sum(1 for x in T if x["zaman"] == "Geciken")}
        elif u["rol"] == "saha":
            T = tum_talepler(d, "WHERE t.sm_adi=? AND t.durum NOT IN ('Tamamlandı','Reddedildi')", (u["sm_adi"],))
            talep = {"liste": T[:6], "acik": len(T), "yeni": sum(1 for x in T if x["durum"] == "Yeni"),
                     "geciken": sum(1 for x in T if x["zaman"] == "Geciken")}
        grafik = {
            "sm": [H.SM_KOD[s] + " " + H.guzel_ad(s.split()[0]) for s in H.SM_ADLARI],
            "tamam": [round(r["tamam_oran"] * 100, 1) for r in sm_tablo],
            "acik": [r["acik"] for r in sm_tablo], "geciken": [r["geciken"] for r in sm_tablo],
            "pp25": [round(pp_sm[s]["pp25"] * 100, 2) for s in H.SM_ADLARI],
            "pp26": [round(pp_sm[s]["pp26"] * 100, 2) for s in H.SM_ADLARI],
        }
        return render_template("pano.html", kpi=kpi, sm_tablo=sm_tablo, oncelikli=oncelikli, isi=isi,
                               hafta=hafta, gecikmis_z=gecikmis_z, benim=benim, grafik=grafik, pp=pp,
                               talep=talep, kredi=kredi, bin_tl=H.bin_tl)

    # ── projeler ────────────────────────────────────────────────────────────
    SIRALAMA = {"id": lambda p: p["id"], "ad": lambda p: p["ad"].lower(), "sm": lambda p: p["sm_adi"],
                "oncelik": lambda p: -H.ONCELIK_AGIRLIK.get(p["oncelik"], 0), "durum": lambda p: p["durum"] or "",
                "termin": lambda p: p["termin"] or "9999", "kalan": lambda p: 99999 if p["kalan_gun"] is None
                else p["kalan_gun"], "skor": lambda p: -p["skor"], "ilerleme": lambda p: p["ilerleme"] or 0,
                "guncelleme": lambda p: p["guncelleme"] or "", "departman": lambda p: p["departman"] or ""}

    @app.route("/projeler")
    @G.giris_gerekli
    def projeler():
        d = vt.baglan()
        f = {k: request.args.get(k, "") for k in ("sm", "dep", "onc", "durum", "zaman", "seviye", "q", "acik", "gorunum")}
        P = tum_projeler(d, f["gorunum"] or "aktif")
        if f["sm"]: P = [p for p in P if p["sm_adi"] == f["sm"]]
        if f["dep"]: P = [p for p in P if p["departman"] == f["dep"]]
        if f["onc"]: P = [p for p in P if p["oncelik"] == f["onc"]]
        if f["durum"]: P = [p for p in P if p["durum"] == f["durum"]]
        if f["zaman"]: P = [p for p in P if p["zaman"] == f["zaman"]]
        if f["seviye"]: P = [p for p in P if p["seviye"] == f["seviye"]]
        if f["acik"] == "1": P = [p for p in P if p["durum"] not in ("Tamamlandı", "İptal")]
        if f["q"]:
            q = f["q"].lower()
            P = [p for p in P if q in (p["ad"] + " " + (p["aciklama"] or "") + " " + p["id"] + " " +
                                       (p["bayi"] or "") + " " + (p["notlar"] or "")).lower()]
        sirala = request.args.get("sirala", "skor")
        yon = request.args.get("yon", "artan")
        P.sort(key=SIRALAMA.get(sirala, SIRALAMA["skor"]), reverse=(yon == "azalan"))
        arsiv_sayi = d.execute("SELECT COUNT(*) FROM projeler WHERE arsiv=1").fetchone()[0]
        return render_template("projeler.html", projeler=P, f=f, sirala=sirala, yon=yon, arsiv_sayi=arsiv_sayi)

    @app.route("/projeler/toplu", methods=["POST"])
    @G.giris_gerekli
    def proje_toplu():
        """Seçilen projeleri topluca arşivle / arşivden çıkar / sil."""
        d = vt.baglan()
        u = G.aktif_kullanici()
        islem = request.form.get("islem")
        secili = list(dict.fromkeys(x for x in request.form.getlist("sec") if x))
        geri = request.form.get("geri") or url_for("projeler")
        if not geri.startswith("/"):
            geri = url_for("projeler")
        if islem not in ("arsivle", "arsivden_cikar", "sil") or not secili:
            flash("Önce en az bir proje seçin.", "hata")
            return redirect(geri)
        satirlar = {r["id"]: r for r in d.execute(
            f"SELECT * FROM projeler WHERE id IN ({','.join('?' * len(secili))})", secili)}
        yetkili, yetkisiz = [], []
        for pid in secili:
            r = satirlar.get(pid)
            if r is None:
                continue
            (yetkili if G.duzenleyebilir(r["sm_adi"]) else yetkisiz).append(pid)
        # ana proje seçildiyse alt adımları da işleme girer
        hedef = list(yetkili)
        for pid in yetkili:
            if not satirlar[pid]["ust_id"]:
                hedef += [x["id"] for x in d.execute("SELECT id FROM projeler WHERE ust_id=?", (pid,))]
        hedef = list(dict.fromkeys(hedef))
        simdi_ = vt.simdi()
        if islem == "arsivle":
            for pid in hedef:
                d.execute("UPDATE projeler SET arsiv=1, arsiv_tarihi=?, arsivleyen=? WHERE id=?",
                          (simdi_, u["kullanici_adi"], pid))
                vt.logla(u["kullanici_adi"], "arşivleme", "proje", pid)
            mesaj = f"{len(hedef)} proje arşive kaldırıldı."
        elif islem == "arsivden_cikar":
            ustler = [satirlar[p]["ust_id"] for p in yetkili if satirlar[p]["ust_id"]]
            for pid in list(dict.fromkeys(hedef + ustler)):   # alt adım geri alınırsa ana projesi de açılır
                d.execute("UPDATE projeler SET arsiv=0, arsiv_tarihi=NULL, arsivleyen=NULL WHERE id=?", (pid,))
                vt.logla(u["kullanici_adi"], "arşivden çıkarma", "proje", pid)
            mesaj = f"{len(hedef)} proje arşivden çıkarıldı."
        else:
            altlar = [p for p in hedef if d.execute("SELECT ust_id FROM projeler WHERE id=?", (p,)).fetchone()["ust_id"]]
            for pid in altlar + [p for p in hedef if p not in altlar]:   # önce alt adımlar
                ad = d.execute("SELECT ad FROM projeler WHERE id=?", (pid,)).fetchone()
                d.execute("UPDATE aksakliklar SET durum='Açık', proje_id=NULL WHERE proje_id=?", (pid,))
                d.execute("UPDATE ziyaretler SET proje_id=NULL WHERE proje_id=?", (pid,))
                d.execute("UPDATE talepler SET proje_id=NULL WHERE proje_id=?", (pid,))
                d.execute("DELETE FROM projeler WHERE id=?", (pid,))
                vt.logla(u["kullanici_adi"], "silme", "proje", pid, ad["ad"] if ad else "")
            mesaj = f"{len(hedef)} proje silindi (gelişme kayıtlarıyla birlikte)."
        d.commit()
        flash(mesaj, "basari")
        if yetkisiz:
            flash(f"{len(yetkisiz)} proje size ait olmadığı için atlandı: {', '.join(yetkisiz[:8])}"
                  f"{'…' if len(yetkisiz) > 8 else ''}", "uyari")
        return redirect(geri)

    def _proje_formu_oku(u, mevcut=None):
        fo = {k: (request.form.get(k) or "").strip() for k in
              ("id", "ust_id", "ad", "aciklama", "sm_adi", "il", "bayi", "departman", "is_konusu", "oncelik",
               "durum", "ilerleme", "baslangic", "termin", "tamamlanma", "notlar")}
        if u["rol"] == "saha":
            fo["sm_adi"] = u["sm_adi"]
        hatalar = []
        if not fo["ad"]: hatalar.append("Proje adı zorunludur.")
        if fo["sm_adi"] not in H.SM_ADLARI: hatalar.append("Geçerli bir saha müdürü seçin.")
        if not fo["termin"]: hatalar.append("Termin tarihi zorunludur — kalan gün ve gecikme buradan hesaplanır.")
        if fo["baslangic"] and fo["termin"] and fo["termin"] < fo["baslangic"]:
            hatalar.append("Termin, başlangıç tarihinden önce olamaz.")
        for alan, liste in (("departman", H.DEPARTMANLAR), ("oncelik", H.ONCELIKLER), ("durum", H.DURUMLAR)):
            if fo[alan] not in liste: hatalar.append(f"Geçersiz {alan} değeri.")
        try:
            il = float(fo["ilerleme"] or 0)
            if not 0 <= il <= 100: raise ValueError
            fo["ilerleme"] = il / 100
        except ValueError:
            hatalar.append("İlerleme 0 ile 100 arasında olmalı."); fo["ilerleme"] = 0
        if fo["durum"] == "Tamamlandı":
            fo["ilerleme"] = 1.0
            fo["tamamlanma"] = fo["tamamlanma"] or H.bugun().isoformat()
        return fo, hatalar

    def _form_secenekleri(d, u, sm=None):
        sm = sm or (u["sm_adi"] if u["rol"] == "saha" else None)
        ana = d.execute("SELECT id, ad, sm_adi FROM projeler WHERE ust_id IS NULL AND COALESCE(arsiv,0)=0 ORDER BY id").fetchall()
        return dict(ana_projeler=[dict(r) for r in ana])

    @app.route("/projeler/yeni", methods=["GET", "POST"])
    @G.giris_gerekli
    def proje_yeni():
        d = vt.baglan()
        u = G.aktif_kullanici()
        if request.method == "POST":
            fo, hatalar = _proje_formu_oku(u)
            ust = None
            if fo["ust_id"]:
                ust = d.execute("SELECT * FROM projeler WHERE id=?", (fo["ust_id"],)).fetchone()
                if not ust: hatalar.append("Seçilen üst proje bulunamadı.")
                elif ust["sm_adi"] != fo["sm_adi"]: hatalar.append("Alt adım, ana projeyle aynı saha müdürüne ait olmalı.")
            if not fo["id"]:
                fo["id"] = H.sonraki_alt_id(d, fo["ust_id"]) if ust else H.sonraki_ana_id(d, fo["sm_adi"])
            if d.execute("SELECT 1 FROM projeler WHERE id=?", (fo["id"],)).fetchone():
                hatalar.append(f"{fo['id']} numaralı proje zaten var.")
            if not H.kod_uyumlu(fo["id"], fo["sm_adi"]):
                hatalar.append(f"Proje ID'si {H.SM_KOD.get(fo['sm_adi'], '?')}- ile başlamalı.")
            G.duzenleme_kontrol(fo["sm_adi"])
            if hatalar:
                for h in hatalar: flash(h, "hata")
                return render_template("proje_form.html", p=fo, yeni=True, **_form_secenekleri(d, u))
            now = vt.simdi()
            d.execute("INSERT INTO projeler (id, ust_id, ad, aciklama, sm_adi, il, bayi, departman, is_konusu, oncelik,"
                      " durum, ilerleme, baslangic, termin, tamamlanma, notlar, olusturan, olusturma, guncelleyen,"
                      " guncelleme) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      (fo["id"], fo["ust_id"] or None, fo["ad"], fo["aciklama"], fo["sm_adi"], fo["il"], fo["bayi"],
                       fo["departman"], fo["is_konusu"], fo["oncelik"], fo["durum"], fo["ilerleme"],
                       fo["baslangic"] or None, fo["termin"], fo["tamamlanma"] or None, fo["notlar"],
                       u["ad_soyad"], now, u["ad_soyad"], now))
            aks = request.form.get("aksaklik_id")
            if aks:
                d.execute("UPDATE aksakliklar SET durum='Projeye dönüştü', proje_id=? WHERE id=?", (fo["id"], aks))
            vt.logla(u["kullanici_adi"], "oluşturma", "proje", fo["id"], fo["ad"])
            d.commit()
            flash(f"{fo['id']} numaralı proje oluşturuldu.", "basari")
            return redirect(url_for("proje_kart", pid=fo["id"]))
        # GET: ön doldurma
        p = {"durum": "Yapılacak", "oncelik": "Orta", "ilerleme": 0, "baslangic": H.bugun().isoformat()}
        if u["rol"] == "saha":
            p["sm_adi"] = u["sm_adi"]
        ust = request.args.get("ust")
        if ust:
            r = d.execute("SELECT * FROM projeler WHERE id=?", (ust,)).fetchone()
            if r:
                G.duzenleme_kontrol(r["sm_adi"])
                p.update(ust_id=r["id"], sm_adi=r["sm_adi"], il=r["il"], bayi=r["bayi"], departman=r["departman"],
                         oncelik=r["oncelik"])
                p["id"] = H.sonraki_alt_id(d, r["id"])
        aks = request.args.get("aksaklik")
        if aks:
            a = d.execute("SELECT * FROM aksakliklar WHERE id=?", (aks,)).fetchone()
            if a:
                G.duzenleme_kontrol(a["sm_adi"])
                onc = {"Kritik": "Kritik", "Yüksek": "Yüksek", "Orta": "Orta"}.get(a["onem"], "Düşük")
                ad = f"{a['bulgu']} – {(a['unvan'] or '')[:40]}"
                p.update(sm_adi=a["sm_adi"], il=(a["il"] or "").upper(), departman=a["departman"], oncelik=onc,
                         ad=ad[:120], aciklama=f"{a['detay']}\nÖnerilen aksiyon: {a['aksiyon']}",
                         aksaklik_id=a["id"],
                         termin=(H.bugun() + dt.timedelta(days=14 if onc == "Kritik" else 30)).isoformat())
        if p.get("sm_adi") and not p.get("id"):
            p["id"] = H.sonraki_ana_id(d, p["sm_adi"])
        return render_template("proje_form.html", p=p, yeni=True, **_form_secenekleri(d, u))

    @app.route("/projeler/<pid>")
    @G.giris_gerekli
    def proje_kart(pid):
        d = vt.baglan()
        r = d.execute("SELECT * FROM projeler WHERE id=?", (pid,)).fetchone()
        if not r:
            abort(404)
        oz = gelisme_ozet(d)
        p = H.proje_zenginlestir(r, H.bugun(), oz)
        ust = d.execute("SELECT id, ad FROM projeler WHERE id=?", (p["ust_id"],)).fetchone() if p["ust_id"] else None
        altlar = [H.proje_zenginlestir(x, H.bugun(), oz) for x in
                  d.execute("SELECT * FROM projeler WHERE ust_id=? ORDER BY id", (pid,))]
        gel = [dict(x) for x in d.execute("SELECT * FROM gelismeler WHERE proje_id=? ORDER BY tarih, id", (pid,))]
        ziy = tum_ziyaretler(d, "WHERE z.proje_id=?", (pid,))
        return render_template("proje_kart.html", p=p, ust=ust, altlar=altlar, gelismeler=gel, ziyaretler=ziy)

    @app.route("/projeler/<pid>/duzenle", methods=["GET", "POST"])
    @G.giris_gerekli
    def proje_duzenle(pid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        r = d.execute("SELECT * FROM projeler WHERE id=?", (pid,)).fetchone()
        if not r:
            abort(404)
        G.duzenleme_kontrol(r["sm_adi"])
        if request.method == "POST":
            fo, hatalar = _proje_formu_oku(u)
            fo["id"], fo["ust_id"] = r["id"], r["ust_id"]
            if u["rol"] == "saha" or r["ust_id"]:
                fo["sm_adi"] = r["sm_adi"]
            if hatalar:
                for h in hatalar: flash(h, "hata")
                return render_template("proje_form.html", p=fo, yeni=False, **_form_secenekleri(d, u))
            degisen = [k for k in ("ad", "durum", "oncelik", "termin", "ilerleme", "departman", "bayi", "il")
                       if str(r[k] or "") != str(fo[k] or "")]
            d.execute("UPDATE projeler SET ad=?, aciklama=?, sm_adi=?, il=?, bayi=?, departman=?, is_konusu=?,"
                      " oncelik=?, durum=?, ilerleme=?, baslangic=?, termin=?, tamamlanma=?, notlar=?,"
                      " guncelleyen=?, guncelleme=? WHERE id=?",
                      (fo["ad"], fo["aciklama"], fo["sm_adi"], fo["il"], fo["bayi"], fo["departman"],
                       fo["is_konusu"], fo["oncelik"], fo["durum"], fo["ilerleme"], fo["baslangic"] or None,
                       fo["termin"], fo["tamamlanma"] or None, fo["notlar"], u["ad_soyad"], vt.simdi(), pid))
            vt.logla(u["kullanici_adi"], "güncelleme", "proje", pid, ", ".join(degisen) or "—")
            d.commit()
            flash("Proje güncellendi.", "basari")
            return redirect(url_for("proje_kart", pid=pid))
        p = dict(r)
        p["ilerleme"] = round((p["ilerleme"] or 0) * 100)
        return render_template("proje_form.html", p=p, yeni=False, **_form_secenekleri(d, u))

    @app.route("/projeler/<pid>/sil", methods=["POST"])
    @G.giris_gerekli
    def proje_sil(pid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        r = d.execute("SELECT * FROM projeler WHERE id=?", (pid,)).fetchone()
        if not r:
            abort(404)
        G.duzenleme_kontrol(r["sm_adi"])
        if d.execute("SELECT 1 FROM projeler WHERE ust_id=?", (pid,)).fetchone():
            flash("Alt adımları olan bir proje silinemez. Önce alt adımları silin.", "hata")
            return redirect(url_for("proje_kart", pid=pid))
        d.execute("UPDATE aksakliklar SET durum='Açık', proje_id=NULL WHERE proje_id=?", (pid,))
        d.execute("UPDATE ziyaretler SET proje_id=NULL WHERE proje_id=?", (pid,))
        d.execute("DELETE FROM projeler WHERE id=?", (pid,))
        vt.logla(u["kullanici_adi"], "silme", "proje", pid, r["ad"])
        d.commit()
        flash(f"{pid} silindi.", "basari")
        return redirect(url_for("projeler"))

    @app.route("/projeler/<pid>/gelisme", methods=["POST"])
    @G.giris_gerekli
    def gelisme_ekle(pid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        r = d.execute("SELECT * FROM projeler WHERE id=?", (pid,)).fetchone()
        if not r:
            abort(404)
        G.duzenleme_kontrol(r["sm_adi"])
        fo = {k: (request.form.get(k) or "").strip() for k in
              ("tarih", "asama", "metin", "ilerleme", "sonraki_adim", "sonraki_tarih")}
        if not fo["metin"]:
            flash("Gelişme metni boş olamaz.", "hata")
            return redirect(url_for("proje_kart", pid=pid) + "#gelisme-ekle")
        if fo["asama"] not in H.ASAMALAR:
            fo["asama"] = "İlerleme"
        tarih_ = fo["tarih"] or H.bugun().isoformat()
        ilr = None
        if fo["ilerleme"]:
            try:
                ilr = max(0, min(100, float(fo["ilerleme"]))) / 100
            except ValueError:
                ilr = None
        d.execute("INSERT INTO gelismeler (proje_id, tarih, asama, metin, ilerleme, sonraki_adim, sonraki_tarih,"
                  " kaydeden, olusturan, olusturma) VALUES (?,?,?,?,?,?,?,?,?,?)",
                  (pid, tarih_, fo["asama"], fo["metin"], ilr, fo["sonraki_adim"], fo["sonraki_tarih"] or None,
                   u["ad_soyad"], u["kullanici_adi"], vt.simdi()))
        guncel = []
        if ilr is not None and request.form.get("ilerleme_yansit"):
            d.execute("UPDATE projeler SET ilerleme=? WHERE id=?", (ilr, pid)); guncel.append("ilerleme")
        if fo["asama"] == "Tamamlandı" and request.form.get("tamamla"):
            d.execute("UPDATE projeler SET durum='Tamamlandı', ilerleme=1, tamamlanma=? WHERE id=?", (tarih_, pid))
            guncel.append("durum=Tamamlandı")
        elif fo["asama"] == "Engel" and r["durum"] == "Devam Ediyor" and request.form.get("beklemeye_al"):
            d.execute("UPDATE projeler SET durum='Beklemede' WHERE id=?", (pid,)); guncel.append("durum=Beklemede")
        elif r["durum"] == "Yapılacak" and fo["asama"] in ("Başlangıç", "İlerleme"):
            d.execute("UPDATE projeler SET durum='Devam Ediyor' WHERE id=?", (pid,)); guncel.append("durum=Devam Ediyor")
        d.execute("UPDATE projeler SET guncelleyen=?, guncelleme=? WHERE id=?", (u["ad_soyad"], vt.simdi(), pid))
        vt.logla(u["kullanici_adi"], "gelişme ekleme", "proje", pid, fo["asama"] + (" · " + ", ".join(guncel) if guncel else ""))
        d.commit()
        flash("Gelişme kaydedildi" + (f" ({', '.join(guncel)})" if guncel else "") + ".", "basari")
        return redirect(url_for("proje_kart", pid=pid) + "#gunluk")

    @app.route("/gelisme/<int:gid>/sil", methods=["POST"])
    @G.giris_gerekli
    def gelisme_sil(gid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        r = d.execute("SELECT g.*, p.sm_adi FROM gelismeler g JOIN projeler p ON p.id=g.proje_id WHERE g.id=?",
                      (gid,)).fetchone()
        if not r:
            abort(404)
        G.duzenleme_kontrol(r["sm_adi"])
        d.execute("DELETE FROM gelismeler WHERE id=?", (gid,))
        vt.logla(u["kullanici_adi"], "gelişme silme", "proje", r["proje_id"], r["metin"][:80])
        d.commit()
        flash("Gelişme kaydı silindi.", "basari")
        return redirect(url_for("proje_kart", pid=r["proje_id"]) + "#gunluk")

    # ── ziyaretler ──────────────────────────────────────────────────────────
    @app.route("/ziyaretler")
    @G.giris_gerekli
    def ziyaretler():
        d = vt.baglan()
        u = G.aktif_kullanici()
        gorunum = request.args.get("gorunum", "hafta")
        sm = request.args.get("sm", "")
        try:
            kay = int(request.args.get("hafta", 0))
        except ValueError:
            kay = 0
        Z = tum_ziyaretler(d)
        if sm:
            Z = [z for z in Z if z["sm_adi"] == sm]
        hb = H.hafta_basi() + dt.timedelta(days=7 * kay)
        gunler = [hb + dt.timedelta(days=i) for i in range(7)]
        izgara = {s: {g_.isoformat(): [] for g_ in gunler} for s in (H.SM_ADLARI if not sm else [sm])}
        for z in Z:
            if z["sm_adi"] in izgara and z["tarih"] in izgara[z["sm_adi"]]:
                izgara[z["sm_adi"]][z["tarih"]].append(z)
        hafta_sonu_dolu = any(izgara[s][g_.isoformat()] for s in izgara for g_ in gunler[5:])
        if not hafta_sonu_dolu:
            gunler = gunler[:5]
        f = {k: request.args.get(k, "") for k in ("durum", "zaman", "amac")}
        L = Z
        if f["durum"]: L = [z for z in L if z["durum"] == f["durum"]]
        if f["zaman"]: L = [z for z in L if z["zaman"] == f["zaman"]]
        if f["amac"]: L = [z for z in L if z["amac"] == f["amac"]]
        L = sorted(L, key=lambda z: (z["tarih"], z["sm_adi"]), reverse=(gorunum == "liste" and not f["zaman"]))
        ozet = {"hafta": sum(1 for z in Z if z["zaman"] == "Bu Hafta"),
                "gecikmis": sum(1 for z in Z if z["zaman"] == "Gecikmiş Plan"),
                "gelecek": sum(1 for z in Z if z["zaman"] == "Gelecek Hafta"),
                "ay_gercek": sum(1 for z in Z if z["durum"] == "Gerçekleşti" and (z["gercek_tarih"] or "") >=
                                 H.bugun().replace(day=1).isoformat())}
        return render_template("ziyaretler.html", gorunum=gorunum, sm=sm, kay=kay, gunler=gunler, izgara=izgara,
                               liste=L, f=f, ozet=ozet, gecikmis=[z for z in Z if z["zaman"] == "Gecikmiş Plan"])

    def _ziyaret_formu_oku(u):
        fo = {k: (request.form.get(k) or "").strip() for k in
              ("tarih", "sm_adi", "bayi_id", "amac", "proje_id", "durum", "gercek_tarih", "notu", "sonraki")}
        if u["rol"] == "saha":
            fo["sm_adi"] = u["sm_adi"]
        h = []
        if not fo["tarih"]: h.append("Planlanan tarih zorunludur.")
        if fo["sm_adi"] not in H.SM_ADLARI: h.append("Geçerli bir saha müdürü seçin.")
        if not fo["bayi_id"]: h.append("Bayi seçin.")
        else:
            b = vt.baglan().execute("SELECT sm_adi FROM bayiler WHERE id=?", (fo["bayi_id"],)).fetchone()
            if not b: h.append("Bayi bulunamadı.")
            elif b["sm_adi"] != fo["sm_adi"]: h.append("Seçilen bayi bu saha müdürünün mıntıkasında değil.")
        if fo["amac"] not in H.ZIYARET_AMACLARI: h.append("Ziyaret amacını seçin.")
        if fo["durum"] not in H.ZIYARET_DURUMLARI: fo["durum"] = "Planlandı"
        if fo["durum"] == "Gerçekleşti" and not fo["gercek_tarih"]:
            fo["gercek_tarih"] = fo["tarih"]
        return fo, h

    def _ziyaret_secenekleri(d):
        return dict(acik_projeler=[dict(r) for r in d.execute(
            "SELECT id, ad, sm_adi FROM projeler WHERE durum NOT IN ('Tamamlandı','İptal') AND COALESCE(arsiv,0)=0 ORDER BY id")])

    @app.route("/ziyaretler/yeni", methods=["GET", "POST"])
    @G.giris_gerekli
    def ziyaret_yeni():
        d = vt.baglan()
        u = G.aktif_kullanici()
        if request.method == "POST":
            fo, h = _ziyaret_formu_oku(u)
            G.duzenleme_kontrol(fo["sm_adi"])
            if h:
                for x in h: flash(x, "hata")
                return render_template("ziyaret_form.html", z=fo, yeni=True, **_ziyaret_secenekleri(d))
            now = vt.simdi()
            cur = d.execute("INSERT INTO ziyaretler (tarih, sm_adi, bayi_id, amac, proje_id, durum, gercek_tarih, notu,"
                            " sonraki, olusturan, olusturma, guncelleyen, guncelleme) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                            (fo["tarih"], fo["sm_adi"], fo["bayi_id"], fo["amac"], fo["proje_id"] or None, fo["durum"],
                             fo["gercek_tarih"] or None, fo["notu"], fo["sonraki"] or None, u["ad_soyad"], now,
                             u["ad_soyad"], now))
            vt.logla(u["kullanici_adi"], "oluşturma", "ziyaret", str(cur.lastrowid), fo["tarih"])
            d.commit()
            flash("Ziyaret planlandı.", "basari")
            return redirect(url_for("ziyaretler", hafta=_hafta_farki(fo["tarih"])))
        z = {"tarih": request.args.get("tarih") or H.bugun().isoformat(), "durum": "Planlandı",
             "sm_adi": u["sm_adi"] if u["rol"] == "saha" else request.args.get("sm", ""),
             "bayi_id": request.args.get("bayi", ""), "proje_id": request.args.get("proje", "")}
        if z["bayi_id"]:
            b = d.execute("SELECT sm_adi FROM bayiler WHERE id=?", (z["bayi_id"],)).fetchone()
            if b: z["sm_adi"] = b["sm_adi"]
        if z["proje_id"]:
            p = d.execute("SELECT sm_adi FROM projeler WHERE id=?", (z["proje_id"],)).fetchone()
            if p: z["sm_adi"] = p["sm_adi"]
        return render_template("ziyaret_form.html", z=z, yeni=True, **_ziyaret_secenekleri(d))

    def _hafta_farki(t):
        d_ = H.tarih(t)
        return ((H.hafta_basi(d_) - H.hafta_basi()).days // 7) if d_ else 0

    @app.route("/ziyaretler/<int:zid>/duzenle", methods=["GET", "POST"])
    @G.giris_gerekli
    def ziyaret_duzenle(zid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        r = d.execute("SELECT * FROM ziyaretler WHERE id=?", (zid,)).fetchone()
        if not r:
            abort(404)
        G.duzenleme_kontrol(r["sm_adi"])
        if request.method == "POST":
            if request.form.get("sil"):
                d.execute("DELETE FROM ziyaretler WHERE id=?", (zid,))
                vt.logla(u["kullanici_adi"], "silme", "ziyaret", str(zid), r["tarih"])
                d.commit()
                flash("Ziyaret silindi.", "basari")
                return redirect(url_for("ziyaretler", hafta=_hafta_farki(r["tarih"])))
            fo, h = _ziyaret_formu_oku(u)
            if u["rol"] == "saha":
                fo["sm_adi"] = r["sm_adi"]
            if h:
                for x in h: flash(x, "hata")
                fo["id"] = zid
                return render_template("ziyaret_form.html", z=fo, yeni=False, **_ziyaret_secenekleri(d))
            d.execute("UPDATE ziyaretler SET tarih=?, sm_adi=?, bayi_id=?, amac=?, proje_id=?, durum=?, gercek_tarih=?,"
                      " notu=?, sonraki=?, guncelleyen=?, guncelleme=? WHERE id=?",
                      (fo["tarih"], fo["sm_adi"], fo["bayi_id"], fo["amac"], fo["proje_id"] or None, fo["durum"],
                       fo["gercek_tarih"] or None, fo["notu"], fo["sonraki"] or None, u["ad_soyad"], vt.simdi(), zid))
            vt.logla(u["kullanici_adi"], "güncelleme", "ziyaret", str(zid), fo["durum"])
            d.commit()
            flash("Ziyaret güncellendi.", "basari")
            return redirect(url_for("ziyaretler", hafta=_hafta_farki(fo["tarih"])))
        return render_template("ziyaret_form.html", z=dict(r), yeni=False, **_ziyaret_secenekleri(d))

    @app.route("/ziyaretler/<int:zid>/tamamla", methods=["POST"])
    @G.giris_gerekli
    def ziyaret_tamamla(zid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        r = d.execute("SELECT * FROM ziyaretler WHERE id=?", (zid,)).fetchone()
        if not r:
            abort(404)
        G.duzenleme_kontrol(r["sm_adi"])
        not_ = (request.form.get("notu") or "").strip()
        d.execute("UPDATE ziyaretler SET durum='Gerçekleşti', gercek_tarih=?, notu=COALESCE(NULLIF(?,''), notu),"
                  " guncelleyen=?, guncelleme=? WHERE id=?", (H.bugun().isoformat(), not_, u["ad_soyad"], vt.simdi(), zid))
        vt.logla(u["kullanici_adi"], "tamamlama", "ziyaret", str(zid))
        d.commit()
        flash("Ziyaret gerçekleşti olarak işaretlendi.", "basari")
        return redirect(request.referrer or url_for("ziyaretler"))

    # ── bayiler ─────────────────────────────────────────────────────────────
    @app.route("/bayiler")
    @G.giris_gerekli
    def bayiler():
        d = vt.baglan()
        risk = bayi_riskleri(d)
        bulgu = {r["bayi_id"]: r["n"] for r in d.execute(
            "SELECT bayi_id, COUNT(*) n FROM aksakliklar WHERE durum='Açık' GROUP BY bayi_id")}
        son_z = {r["bayi_id"]: r["t"] for r in d.execute(
            "SELECT bayi_id, MAX(COALESCE(gercek_tarih, tarih)) t FROM ziyaretler WHERE durum='Gerçekleşti' GROUP BY bayi_id")}
        B = []
        for r in d.execute("SELECT * FROM bayiler"):
            b = dict(r)
            b["risk"] = risk.get(b["id"], "Temiz")
            b["bulgu"] = bulgu.get(b["id"], 0)
            b["yoy"] = (b["ytd26_hacim"] / b["ytd25_hacim"] - 1) if b["ytd25_hacim"] else None
            b["trend"] = (b["h2"] / b["h1"] - 1) if b["h1"] else None
            b["kar_m3"] = (b["ytd26_kar"] / b["ytd26_hacim"]) if b["ytd26_hacim"] else None
            b["son_ziyaret"] = son_z.get(b["id"])
            B.append(b)
        f = {k: request.args.get(k, "") for k in ("sm", "il", "segment", "risk", "q")}
        if f["sm"]: B = [b for b in B if b["sm_adi"] == f["sm"]]
        if f["il"]: B = [b for b in B if b["il"] == f["il"]]
        if f["segment"]: B = [b for b in B if b["segment"] == f["segment"]]
        if f["risk"]: B = [b for b in B if b["risk"] == f["risk"]]
        if f["q"]:
            q = f["q"].lower()
            B = [b for b in B if q in f"{b['unvan']} {b['ist_kodu']} {b['ilce']}".lower()]
        sirala = request.args.get("sirala", "hacim")
        anahtar = {"hacim": lambda b: -(b["ay_hacim"] or 0), "unvan": lambda b: b["unvan"],
                   "yoy": lambda b: b["yoy"] if b["yoy"] is not None else 9, "kar": lambda b: b["ytd26_kar"] or 0,
                   "risk": lambda b: H.ONEM_SIRA.get(b["risk"], 9), "sm": lambda b: (b["sm_adi"], -(b["ay_hacim"] or 0)),
                   "ziyaret": lambda b: b["son_ziyaret"] or "0000"}
        B.sort(key=anahtar.get(sirala, anahtar["hacim"]))
        iller = sorted({r["il"] for r in d.execute("SELECT DISTINCT il FROM bayiler")})
        return render_template("bayiler.html", bayiler=B, f=f, iller=iller, sirala=sirala)

    @app.route("/bayiler/<int:bid>")
    @G.giris_gerekli
    def bayi_kart(bid):
        d = vt.baglan()
        r = d.execute("SELECT * FROM bayiler WHERE id=?", (bid,)).fetchone()
        if not r:
            abort(404)
        b = dict(r)
        b["risk"] = bayi_riskleri(d).get(bid, "Temiz")
        b["yoy"] = (b["ytd26_hacim"] / b["ytd25_hacim"] - 1) if b["ytd25_hacim"] else None
        b["trend"] = (b["h2"] / b["h1"] - 1) if b["h1"] else None
        b["kar_m3"] = (b["ytd26_kar"] / b["ytd26_hacim"]) if b["ytd26_hacim"] else None
        aks = [dict(x) for x in d.execute("SELECT * FROM aksakliklar WHERE bayi_id=? ORDER BY durum, onem", (bid,))]
        aks.sort(key=lambda a: (a["durum"] != "Açık", H.ONEM_SIRA.get(a["onem"], 9)))
        ziy = tum_ziyaretler(d, "WHERE z.bayi_id=?", (bid,))
        br = d.execute("SELECT * FROM borc WHERE bayi_id=?", (bid,)).fetchone()
        borc = H.borc_zenginlestir(br) if br else None
        onaylar = [dict(x) for x in d.execute("SELECT * FROM risk_onaylari WHERE bayi_id=? ORDER BY zaman DESC"
                                              " LIMIT 15", (bid,))]
        return render_template("bayi_kart.html", b=b, aksakliklar=aks, ziyaretler=list(reversed(ziy)),
                               borc=borc, onaylar=onaylar, bin_tl=H.bin_tl)

    # ── aksaklıklar ─────────────────────────────────────────────────────────
    @app.route("/aksakliklar")
    @G.giris_gerekli
    def aksakliklar():
        d = vt.baglan()
        A = [dict(r) for r in d.execute("SELECT * FROM aksakliklar")]
        f = {k: request.args.get(k, "") for k in ("sm", "onem", "kategori", "durum", "q")}
        if "durum" not in request.args:
            f["durum"] = "Açık"
        if f["sm"]: A = [a for a in A if a["sm_adi"] == f["sm"]]
        if f["onem"]: A = [a for a in A if a["onem"] == f["onem"]]
        if f["kategori"]: A = [a for a in A if a["kategori"] == f["kategori"]]
        if f["durum"]: A = [a for a in A if a["durum"] == f["durum"]]
        if f["q"]:
            q = f["q"].lower()
            A = [a for a in A if q in f"{a['unvan']} {a['bulgu']} {a['detay']} {a['ilce']}".lower()]
        A.sort(key=lambda a: (H.ONEM_SIRA.get(a["onem"], 9), a["sm_adi"] or "", a["kategori"] or ""))
        kategoriler = sorted({r[0] for r in d.execute("SELECT DISTINCT kategori FROM aksakliklar")})
        sayim = {r["durum"]: r["n"] for r in d.execute("SELECT durum, COUNT(*) n FROM aksakliklar GROUP BY durum")}
        return render_template("aksakliklar.html", aksakliklar=A, f=f, kategoriler=kategoriler, sayim=sayim)

    @app.route("/aksakliklar/<int:aid>/durum", methods=["POST"])
    @G.giris_gerekli
    def aksaklik_durum(aid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        r = d.execute("SELECT * FROM aksakliklar WHERE id=?", (aid,)).fetchone()
        if not r:
            abort(404)
        G.duzenleme_kontrol(r["sm_adi"])
        yeni = request.form.get("durum")
        if yeni not in H.AKSAKLIK_DURUMLARI:
            abort(400)
        d.execute("UPDATE aksakliklar SET durum=? WHERE id=?", (yeni, aid))
        vt.logla(u["kullanici_adi"], "durum değişikliği", "aksaklık", str(aid), yeni)
        d.commit()
        flash(f"Bulgu durumu: {yeni}", "basari")
        return redirect(request.referrer or url_for("aksakliklar"))

    # ── analiz sayfaları ────────────────────────────────────────────────────
    @app.route("/mintika")
    @G.giris_gerekli
    def mintika():
        return render_template("mintika.html", m=mintika_hesapla(vt.baglan()))

    @app.route("/pazar-payi")
    @G.giris_gerekli
    def pazar_payi():
        pp = pazar_payi_hesapla(vt.baglan())
        grafik = {"iller": [r["il"] for r in pp["sehir"]], "bp": [round(r["bp"]) for r in pp["sehir"]],
                  "sm": [H.guzel_ad(s.split()[0]) for s in H.SM_ADLARI],
                  "pp25": [round(r["pp25"] * 100, 2) for r in pp["sm"]],
                  "pp26": [round(r["pp26"] * 100, 2) for r in pp["sm"]],
                  "aylar": [r["ay"] for r in pp["aylik"]],
                  "a25": [round(r["pp25"] * 100, 2) for r in pp["aylik"]],
                  "a26": [round(r["pp26"] * 100, 2) for r in pp["aylik"]]}
        return render_template("pazar_payi.html", pp=pp, grafik=grafik)

    @app.route("/kilavuz")
    @G.giris_gerekli
    def kilavuz():
        return render_template("kilavuz.html")

    # ── dışa aktarım ────────────────────────────────────────────────────────
    @app.route("/disa-aktar")
    @G.giris_gerekli
    def disa_aktar():
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        d = vt.baglan()
        u = G.aktif_kullanici()
        wb = Workbook()
        def sayfa(ad, basliklar, satirlar, ilk=False):
            ws = wb.active if ilk else wb.create_sheet(ad)
            ws.title = ad
            ws.append(basliklar)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for s in satirlar:
                ws.append(list(s))
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or "")) for c in col[:60]) + 2))
        P = tum_projeler(d, "tumu")
        P.sort(key=lambda p: p["id"])
        sayfa("Projeler", ["Proje ID", "Üst Proje ID", "Seviye", "Proje Adı", "Açıklama", "Saha Müdürü", "İl", "Bayi",
                           "Departman", "İş Konusu", "Öncelik", "Durum", "İlerleme %", "Başlangıç", "Termin",
                           "Tamamlanma", "Kalan Gün", "Gecikme", "Zaman Durumu", "Gelişme Sayısı", "Son Gelişme Tarihi",
                           "Son Gelişme", "Notlar", "Son Güncelleyen", "Son Güncelleme", "Arşiv"],
              [(p["id"], p["ust_id"], p["seviye"], p["ad"], p["aciklama"], p["sm_adi"], p["il"], p["bayi"],
                p["departman"], p["is_konusu"], p["oncelik"], p["durum"], round((p["ilerleme"] or 0) * 100),
                p["baslangic"], p["termin"], p["tamamlanma"], p["kalan_gun"], p["gecikme"], p["zaman"],
                p["gelisme_sayisi"], p["son_gelisme_tarihi"], p["son_gelisme"], p["notlar"], p["guncelleyen"],
                p["guncelleme"], "Arşivde" if p.get("arsiv") else "") for p in P], ilk=True)
        sayfa("Gelişmeler", ["Tarih", "Proje ID", "Proje Adı", "Saha Müdürü", "Aşama", "Gelişme", "İlerleme %",
                             "Sonraki Adım", "Sonraki Tarih", "Kaydeden"],
              [(r["tarih"], r["proje_id"], r["ad"], r["sm_adi"], r["asama"], r["metin"],
                round(r["ilerleme"] * 100) if r["ilerleme"] is not None else None, r["sonraki_adim"],
                r["sonraki_tarih"], r["kaydeden"]) for r in d.execute(
                  "SELECT g.*, p.ad, p.sm_adi FROM gelismeler g JOIN projeler p ON p.id=g.proje_id ORDER BY g.tarih, g.id")])
        sayfa("Ziyaretler", ["Tarih", "Gün", "Saha Müdürü", "Bayi", "İst. Kodu", "İl", "İlçe", "Segment", "Amaç",
                             "İlgili Proje", "Durum", "Zaman Durumu", "Gerçekleşen", "Not", "Sonraki Ziyaret"],
              [(z["tarih"], H.GUNLER[H.tarih(z["tarih"]).weekday()], z["sm_adi"], z["unvan"], z["ist_kodu"], z["il"],
                z["ilce"], z["segment"], z["amac"], z["proje_id"], z["durum"], z["zaman"], z["gercek_tarih"],
                z["notu"], z["sonraki"]) for z in tum_ziyaretler(d)])
        sayfa("Aksaklıklar", ["Önem", "Kategori", "Durum", "Saha Müdürü", "İl", "İlçe", "Bayi", "Bulgu", "Detay",
                              "Departman", "Önerilen Aksiyon", "Proje ID"],
              [(a["onem"], a["kategori"], a["durum"], a["sm_adi"], a["il"], a["ilce"], a["unvan"], a["bulgu"],
                a["detay"], a["departman"], a["aksiyon"], a["proje_id"]) for a in d.execute("SELECT * FROM aksakliklar")])
        BR = [H.borc_zenginlestir(r) for r in d.execute("SELECT * FROM borc ORDER BY vadesi_gecmis DESC")]
        sayfa("Kredi-Risk (Borç)", ["SAP", "Bayi", "İl", "Marka", "Saha Müdürü", "Önceki (17.09)", "Fark",
                                    "Vadesi Geçmiş", "1-7 Gün", "8-15 Gün", "16-30 Gün", "31-45 Gün", "45 Gün+",
                                    "Açık Hesap", "Vadesi Gelmemiş", "Çek/Senet", "Toplam Risk", "Seviye",
                                    "Risk Skoru", "Açıklama"],
              [(b["sap"], b["unvan"], b["il"], b["marka"], b["sm_adi"], b["onceki"], b["fark"], b["vadesi_gecmis"],
                b["g1_7"], b["g8_15"], b["g16_30"], b["g31_45"], b["g45_ustu"], b["acik_hesap"],
                b["vadesi_gelmemis"], b["cek_senet"], b["toplam_risk"], b["seviye"], b["skor"], b["aciklama"])
               for b in BR])
        sayfa("Kredi-Risk (Onaylar)", ["Tarih", "Saha Müdürü", "Bayi", "Tür", "Unpaid (bin TL)", "Limit (bin TL)",
                                       "Risk (bin TL)", "Sipariş (bin TL)", "İpotek", "BTM", "Durum",
                                       "Yanıt süresi (dk)", "Talep", "Yanıt"],
              [(o["zaman"].replace("T", " "), o["sm_adi"], o["bayi"], o["tur"], o["unpaid"], o["limit_tl"],
                o["risk"], o["siparis"], o["ipotek"], o["btm"], o["onay_durumu"], o["sure_dk"],
                (o["metin"] or "")[:800], (o["yanit"] or "")[:300])
               for o in d.execute("SELECT * FROM risk_onaylari ORDER BY zaman")])
        TL = tum_talepler(d) if u["rol"] == "admin" else tum_talepler(d, "WHERE t.sm_adi=?", (u["sm_adi"],))
        sayfa("BM'ye Atanan İşler", ["No", "Saha Müdürü", "İş", "Açıklama", "Departman", "Öncelik", "İstenen Tarih",
                                     "Kalan Gün", "Durum", "Bölge Müdürü Yanıtı", "Bayi", "İlgili Proje", "Açılış",
                                     "Tamamlanma"],
              [(t["id"], t["sm_adi"], t["baslik"], t["aciklama"], t["departman"], t["oncelik"], t["termin"],
                t["kalan_gun"], t["durum"], t["yanit"], t["unvan"], t["proje_id"], (t["olusturma"] or "")[:10],
                t["tamamlanma"]) for t in TL])
        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        vt.logla(u["kullanici_adi"], "dışa aktarım", "excel"); d.commit()
        return send_file(buf, as_attachment=True, download_name=f"Marmara_Bolge_{H.bugun().isoformat()}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.route("/kredi-risk")
    @G.giris_gerekli
    def kredi_risk():
        d = vt.baglan()
        f = {k: request.args.get(k, "") for k in ("sm", "seviye", "q", "durum", "gorunum")}
        K = kredi_hesapla(d)
        B = K["borclu"]
        if f["sm"]: B = [b for b in B if b["sm_adi"] == f["sm"]]
        if f["seviye"]: B = [b for b in B if b["seviye"] == f["seviye"]]
        if f["q"]:
            q = f["q"].lower()
            B = [b for b in B if q in ((b["unvan"] or "") + " " + (b["il"] or "") + " " +
                                       (b["aciklama"] or "")).lower()]
        O = K["onaylar"]
        if f["sm"]: O = [o for o in O if o["sm_adi"] == f["sm"]]
        if f["durum"]: O = [o for o in O if o["onay_durumu"] == f["durum"]]
        if f["q"]:
            q = f["q"].lower()
            O = [o for o in O if q in ((o["bayi"] or "") + " " + (o["metin"] or "")).lower()]
        grafik = {
            "sm": [x["kod"] for x in K["sm"]],
            "yas": [{"ad": ad, "veri": [round((x[k] or 0) / 1000) for x in K["sm"]]} for k, ad in H.YAS_KOLONLARI],
            "onay": [x["onay"] for x in K["sm"]],
        }
        return render_template("kredi_risk.html", k=K, borclu=B[:60], onaylar=O[:120], onay_sayi=len(O),
                               f=f, grafik=grafik, ay=H.bin_tl)

    # ── bölge müdürüne atanan işler (saha → BM) ─────────────────────────────
    def _talep_yetki(t, yazma=False):
        """Bölge müdürü tüm işleri, saha müdürü yalnızca kendi atadıklarını görür/yönetir."""
        u = G.aktif_kullanici()
        if u["rol"] == "admin":
            return
        if u["rol"] != "saha" or t["sm_adi"] != u["sm_adi"]:
            abort(403)

    def _talep_formu(u):
        fo = {k: (request.form.get(k) or "").strip() for k in
              ("baslik", "aciklama", "departman", "oncelik", "termin", "bayi_id", "proje_id")}
        hatalar = []
        if len(fo["baslik"]) < 3: hatalar.append("İşin başlığını yazın.")
        if fo["oncelik"] not in H.ONCELIKLER: hatalar.append("Öncelik seçin.")
        if fo["departman"] and fo["departman"] not in H.DEPARTMANLAR: hatalar.append("Geçersiz departman.")
        if fo["termin"]:
            try: dt.date.fromisoformat(fo["termin"])
            except ValueError: hatalar.append("Tarih geçersiz.")
        d = vt.baglan()
        fo["bayi_id"] = int(fo["bayi_id"]) if fo["bayi_id"].isdigit() else None
        if fo["bayi_id"] and not d.execute("SELECT 1 FROM bayiler WHERE id=? AND sm_adi=?",
                                           (fo["bayi_id"], u["sm_adi"])).fetchone():
            hatalar.append("Bayi yalnızca kendi mıntıkanızdan seçilebilir.")
        if fo["proje_id"] and not d.execute("SELECT 1 FROM projeler WHERE id=? AND sm_adi=?",
                                            (fo["proje_id"], u["sm_adi"])).fetchone():
            hatalar.append("İlgili proje yalnızca kendi projelerinizden seçilebilir.")
        return fo, hatalar

    def _talep_secenekleri(d, u):
        return dict(
            bayiler=[dict(r) for r in d.execute("SELECT id, unvan, il, ilce FROM bayiler WHERE sm_adi=? ORDER BY il, unvan",
                                                (u["sm_adi"],))],
            projeler_=[dict(r) for r in d.execute("SELECT id, ad FROM projeler WHERE sm_adi=? AND COALESCE(arsiv,0)=0"
                                                  " AND durum NOT IN ('Tamamlandı','İptal') ORDER BY id", (u["sm_adi"],))])

    @app.route("/isler")
    @G.giris_gerekli
    def talepler():
        u = G.aktif_kullanici()
        if u["rol"] == "ziyaretci":
            abort(403)
        d = vt.baglan()
        f = {k: request.args.get(k, "") for k in ("sm", "durum", "kapali")}
        where, prm = [], []
        if u["rol"] == "saha":
            where.append("t.sm_adi=?"); prm.append(u["sm_adi"])
        elif f["sm"]:
            where.append("t.sm_adi=?"); prm.append(f["sm"])
        if f["durum"]:
            where.append("t.durum=?"); prm.append(f["durum"])
        elif f["kapali"] != "1":
            where.append("t.durum NOT IN ('Tamamlandı','Reddedildi')")
        T = tum_talepler(d, ("WHERE " + " AND ".join(where)) if where else "", prm)
        ozet = {"yeni": sum(1 for t in T if t["durum"] == "Yeni"),
                "acik": sum(1 for t in T if t["durum"] not in H.TALEP_KAPALI),
                "geciken": sum(1 for t in T if t["zaman"] == "Geciken")}
        return render_template("talepler.html", talepler=T, f=f, ozet=ozet)

    @app.route("/isler/yeni", methods=["GET", "POST"])
    @G.giris_gerekli
    def talep_yeni():
        u = G.aktif_kullanici()
        if u["rol"] != "saha":
            abort(403)
        d = vt.baglan()
        fo = {"baslik": "", "aciklama": "", "departman": "", "oncelik": "Orta", "termin": "",
              "bayi_id": request.args.get("bayi", type=int), "proje_id": request.args.get("proje", "")}
        if request.method == "POST":
            fo, hatalar = _talep_formu(u)
            if hatalar:
                for h in hatalar: flash(h, "hata")
            else:
                cur = d.execute("INSERT INTO talepler (sm_adi, olusturan, baslik, aciklama, departman, oncelik, termin,"
                                " bayi_id, proje_id, durum, olusturma) VALUES (?,?,?,?,?,?,?,?,?, 'Yeni', ?)",
                                (u["sm_adi"], u["kullanici_adi"], fo["baslik"], fo["aciklama"], fo["departman"] or None,
                                 fo["oncelik"], fo["termin"] or None, fo["bayi_id"], fo["proje_id"] or None, vt.simdi()))
                vt.logla(u["kullanici_adi"], "iş atama", "bm işi", str(cur.lastrowid), fo["baslik"])
                d.commit()
                flash("İş bölge müdürüne iletildi.", "basari")
                return redirect(url_for("talep_kart", tid=cur.lastrowid))
        return render_template("talep_form.html", t=fo, yeni=True, **_talep_secenekleri(d, u))

    @app.route("/isler/<int:tid>")
    @G.giris_gerekli
    def talep_kart(tid):
        d = vt.baglan()
        r = tum_talepler(d, "WHERE t.id=?", (tid,))
        if not r:
            abort(404)
        t = r[0]
        _talep_yetki(t)
        notlar = [dict(x) for x in d.execute("SELECT * FROM talep_notlari WHERE talep_id=? ORDER BY id", (tid,))]
        proje = d.execute("SELECT id, ad FROM projeler WHERE id=?", (t["proje_id"],)).fetchone() if t["proje_id"] else None
        return render_template("talep_kart.html", t=t, notlar=notlar, proje=proje)

    @app.route("/isler/<int:tid>/duzenle", methods=["GET", "POST"])
    @G.giris_gerekli
    def talep_duzenle(tid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        t = d.execute("SELECT * FROM talepler WHERE id=?", (tid,)).fetchone()
        if not t:
            abort(404)
        if u["rol"] != "saha" or t["sm_adi"] != u["sm_adi"]:
            abort(403)
        if t["durum"] in H.TALEP_KAPALI:
            flash("Kapanmış bir iş düzenlenemez; yeni bir iş açın.", "hata")
            return redirect(url_for("talep_kart", tid=tid))
        fo = dict(t)
        if request.method == "POST":
            if request.form.get("geri_cek"):
                d.execute("DELETE FROM talepler WHERE id=?", (tid,))
                vt.logla(u["kullanici_adi"], "iş geri çekme", "bm işi", str(tid), t["baslik"])
                d.commit()
                flash("İş geri çekildi.", "bilgi")
                return redirect(url_for("talepler"))
            fo, hatalar = _talep_formu(u)
            if hatalar:
                for h in hatalar: flash(h, "hata")
            else:
                d.execute("UPDATE talepler SET baslik=?, aciklama=?, departman=?, oncelik=?, termin=?, bayi_id=?,"
                          " proje_id=?, guncelleyen=?, guncelleme=? WHERE id=?",
                          (fo["baslik"], fo["aciklama"], fo["departman"] or None, fo["oncelik"], fo["termin"] or None,
                           fo["bayi_id"], fo["proje_id"] or None, u["kullanici_adi"], vt.simdi(), tid))
                vt.logla(u["kullanici_adi"], "iş güncelleme", "bm işi", str(tid))
                d.commit()
                flash("İş güncellendi.", "basari")
                return redirect(url_for("talep_kart", tid=tid))
        fo["id"] = tid
        return render_template("talep_form.html", t=fo, yeni=False, **_talep_secenekleri(d, u))

    @app.route("/isler/<int:tid>/durum", methods=["POST"])
    @G.admin_gerekli
    def talep_durum(tid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        t = d.execute("SELECT * FROM talepler WHERE id=?", (tid,)).fetchone()
        if not t:
            abort(404)
        durum = request.form.get("durum")
        yanit = (request.form.get("yanit") or "").strip()
        if durum not in H.TALEP_DURUMLARI:
            abort(400)
        d.execute("UPDATE talepler SET durum=?, yanit=?, tamamlanma=?, guncelleyen=?, guncelleme=? WHERE id=?",
                  (durum, yanit or t["yanit"], H.bugun().isoformat() if durum == "Tamamlandı" else None,
                   u["kullanici_adi"], vt.simdi(), tid))
        if durum != t["durum"] or (yanit and yanit != (t["yanit"] or "")):
            d.execute("INSERT INTO talep_notlari (talep_id, yazan, yazan_ad, rol, metin, zaman) VALUES (?,?,?,?,?,?)",
                      (tid, u["kullanici_adi"], u["ad_soyad"], u["rol"],
                       f"Durum: {durum}" + (f" — {yanit}" if yanit else ""), vt.simdi()))
        vt.logla(u["kullanici_adi"], "iş durumu", "bm işi", str(tid), durum)
        d.commit()
        flash(f"İş #{tid}: {durum}", "basari")
        return redirect(request.form.get("geri") if (request.form.get("geri") or "").startswith("/")
                        else url_for("talep_kart", tid=tid))

    @app.route("/isler/<int:tid>/not", methods=["POST"])
    @G.giris_gerekli
    def talep_not(tid):
        d = vt.baglan()
        u = G.aktif_kullanici()
        t = d.execute("SELECT * FROM talepler WHERE id=?", (tid,)).fetchone()
        if not t:
            abort(404)
        _talep_yetki(t)
        metin = (request.form.get("metin") or "").strip()
        if metin:
            d.execute("INSERT INTO talep_notlari (talep_id, yazan, yazan_ad, rol, metin, zaman) VALUES (?,?,?,?,?,?)",
                      (tid, u["kullanici_adi"], u["ad_soyad"], u["rol"], metin[:2000], vt.simdi()))
            if u["rol"] == "admin" and t["durum"] == "Yeni":
                d.execute("UPDATE talepler SET durum='İnceleniyor', guncelleme=? WHERE id=?", (vt.simdi(), tid))
            vt.logla(u["kullanici_adi"], "iş notu", "bm işi", str(tid))
            d.commit()
        return redirect(url_for("talep_kart", tid=tid) + "#notlar")

    # ── yönetim ─────────────────────────────────────────────────────────────
    @app.route("/yonetim/kullanicilar", methods=["GET", "POST"])
    @G.admin_gerekli
    def kullanicilar():
        d = vt.baglan()
        u = G.aktif_kullanici()
        yeni_sifre = None
        if request.method == "POST":
            islem = request.form.get("islem")
            if islem == "yeni":
                ka = request.form.get("kullanici_adi", "").strip().lower()
                ad = request.form.get("ad_soyad", "").strip()
                rol = request.form.get("rol")
                sm = request.form.get("sm_adi") or None
                if not ka.isalnum() or not ad or rol not in ("admin", "saha", "ziyaretci") or (rol == "saha" and sm not in H.SM_ADLARI):
                    flash("Kullanıcı bilgilerini kontrol edin (kullanıcı adı yalnızca harf/rakam).", "hata")
                elif d.execute("SELECT 1 FROM kullanicilar WHERE kullanici_adi=?", (ka,)).fetchone():
                    flash("Bu kullanıcı adı zaten var.", "hata")
                else:
                    s = vt.rastgele_sifre()
                    d.execute("INSERT INTO kullanicilar (kullanici_adi, ad_soyad, rol, sm_adi, sm_kodu, sifre_hash,"
                              " sifre_degismeli, olusturma) VALUES (?,?,?,?,?,?,1,?)",
                              (ka, ad, rol, sm if rol == "saha" else None,
                               H.SM_KOD.get(sm, "NW1") if rol == "saha" else ("NW1" if rol == "admin" else None),
                               generate_password_hash(s), vt.simdi()))
                    vt.logla(u["kullanici_adi"], "kullanıcı oluşturma", "kullanıcı", ka)
                    d.commit()
                    yeni_sifre = (ka, s)
            else:
                kid = request.form.get("id", type=int)
                k = d.execute("SELECT * FROM kullanicilar WHERE id=?", (kid,)).fetchone()
                if not k:
                    abort(404)
                if islem == "onayla":
                    mk = request.form.get("mintika")
                    rol = request.form.get("rol", "saha")
                    kodlar = {k: s for s, k in H.SAHA_MUDURLERI}
                    if rol not in ROL_ADI or (rol == "saha" and mk not in kodlar):
                        flash("Onay için geçerli bir rol ve mıntıka seçin.", "hata")
                    else:
                        sm_kodu = mk if rol == "saha" else ("NW1" if rol == "admin" else None)
                        d.execute("UPDATE kullanicilar SET aktif=1, onay_durumu='onaylı', rol=?, sm_adi=?, sm_kodu=? WHERE id=?",
                                  (rol, kodlar.get(mk) if rol == "saha" else None, sm_kodu, kid))
                        vt.logla(u["kullanici_adi"], "kayıt onayı", "kullanıcı", k["kullanici_adi"],
                                 mk if rol == "saha" else ROL_ADI[rol])
                        d.commit()
                        flash(f"{k['ad_soyad']} onaylandı — artık giriş yapabilir.", "basari")
                elif islem == "reddet":
                    d.execute("UPDATE kullanicilar SET aktif=0, onay_durumu='reddedildi' WHERE id=?", (kid,))
                    vt.logla(u["kullanici_adi"], "kayıt reddi", "kullanıcı", k["kullanici_adi"])
                    d.commit()
                    flash(f"{k['ad_soyad']} başvurusu reddedildi.", "bilgi")
                elif islem == "sifirla":
                    s = vt.rastgele_sifre()
                    d.execute("UPDATE kullanicilar SET sifre_hash=?, sifre_degismeli=1, hatali_giris=0, kilit_bitis=NULL"
                              " WHERE id=?", (generate_password_hash(s), kid))
                    vt.logla(u["kullanici_adi"], "şifre sıfırlama", "kullanıcı", k["kullanici_adi"])
                    d.commit()
                    yeni_sifre = (k["kullanici_adi"], s)
                elif islem == "aktiflik":
                    if k["id"] == u["id"]:
                        flash("Kendi hesabınızı pasif yapamazsınız.", "hata")
                    else:
                        d.execute("UPDATE kullanicilar SET aktif=1-aktif, onay_durumu=CASE WHEN aktif=0 THEN 'onaylı'"
                                  " ELSE onay_durumu END WHERE id=?", (kid,))
                        vt.logla(u["kullanici_adi"], "aktiflik değişikliği", "kullanıcı", k["kullanici_adi"])
                        d.commit()
        K = [dict(r) for r in d.execute("SELECT * FROM kullanicilar WHERE onay_durumu!='onay_bekliyor'"
                                        " ORDER BY rol, sm_kodu, kullanici_adi")]
        B = [dict(r) for r in d.execute("SELECT * FROM kullanicilar WHERE onay_durumu='onay_bekliyor' ORDER BY id")]
        mevcut = defaultdict(list)
        for r in d.execute("SELECT sm_kodu, ad_soyad FROM kullanicilar WHERE aktif=1 AND rol='saha'"):
            mevcut[r["sm_kodu"]].append(r["ad_soyad"])
        return render_template("kullanicilar.html", kullanicilar=K, bekleyenler=B, mevcut=mevcut, yeni_sifre=yeni_sifre)

    @app.route("/yonetim/islem-gecmisi")
    @G.admin_gerekli
    def islem_gecmisi():
        L = [dict(r) for r in vt.baglan().execute("SELECT * FROM islem_gecmisi ORDER BY id DESC LIMIT 400")]
        return render_template("islem_gecmisi.html", kayitlar=L)

    @app.route("/yonetim/yedek")
    @G.admin_gerekli
    def yedek():
        u = G.aktif_kullanici()
        tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False); tmp.close()
        kaynak = sqlite3.connect(app.config["VERITABANI"])
        hedef = sqlite3.connect(tmp.name)
        kaynak.backup(hedef); hedef.close(); kaynak.close()
        vt.logla(u["kullanici_adi"], "yedek alma", "veritabanı"); vt.baglan().commit()
        return send_file(tmp.name, as_attachment=True,
                         download_name=f"platform_yedek_{dt.datetime.now():%Y%m%d_%H%M}.db")

    # ── API (kademeli listeler) ─────────────────────────────────────────────
    @app.route("/api/kanopiler")
    @G.giris_gerekli
    def api_kanopiler():
        sm = request.args.get("sm", "")
        rows = vt.baglan().execute("SELECT etiket, il, ilce FROM kanopiler WHERE sm_adi=? ORDER BY etiket", (sm,))
        return jsonify([dict(r) for r in rows])

    @app.route("/api/bayiler")
    @G.giris_gerekli
    def api_bayiler():
        d = vt.baglan()
        sm = request.args.get("sm", "")
        risk = bayi_riskleri(d)
        rows = d.execute("SELECT id, unvan, ist_kodu, il, ilce, segment, ay_hacim FROM bayiler WHERE sm_adi=?"
                         " ORDER BY il, ilce, unvan", (sm,))
        return jsonify([dict(r, risk=risk.get(r["id"], "Temiz")) for r in rows])

    @app.route("/api/sonraki-id")
    @G.giris_gerekli
    def api_sonraki_id():
        d = vt.baglan()
        ust = request.args.get("ust", "")
        if ust:
            return jsonify({"id": H.sonraki_alt_id(d, ust)})
        return jsonify({"id": H.sonraki_ana_id(d, request.args.get("sm", ""))})

    @app.route("/api/iller")
    @G.giris_gerekli
    def api_iller():
        sm = request.args.get("sm", "")
        rows = vt.baglan().execute("SELECT DISTINCT il FROM kanopiler WHERE sm_adi=? ORDER BY il", (sm,))
        return jsonify([r["il"] for r in rows])


def _bos_port(host, baslangic):
    """İstenen port doluysa sıradaki boş portu bulur (ör. 5000 başka program tarafından kullanılıyorsa 5001)."""
    import socket
    for p in range(baslangic, baslangic + 20):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            try:
                s.bind((host, p))
                return p
            except OSError:
                continue
    return baslangic


def _hazir_olunca_tarayici_ac(port):
    """Sunucu gerçekten bağlantı kabul etmeye başlayınca tarayıcıyı açar."""
    import socket, threading, time, webbrowser
    def calis():
        for _ in range(120):
            try:
                socket.create_connection(("127.0.0.1", port), timeout=0.5).close()
                webbrowser.open(f"http://127.0.0.1:{port}")
                return
            except OSError:
                time.sleep(0.5)
    threading.Thread(target=calis, daemon=True).start()


if __name__ == "__main__":
    print("\n  Platform hazırlanıyor...")
    app = uygulama_olustur()
    host = os.environ.get("HOST", "127.0.0.1")
    istenen = int(os.environ.get("PORT", 5000))
    port = _bos_port(host, istenen)
    if port != istenen:
        print(f"  Not: {istenen} numaralı port başka bir program tarafından kullanılıyor, {port} kullanılacak.")
    print("\n  " + "-" * 60)
    print(f"  PLATFORM ÇALIŞIYOR  →  http://127.0.0.1:{port}")
    print("  Bu pencere açık kaldığı sürece platform çalışır.")
    print("  Kapatmak için bu pencereyi kapatın veya Ctrl+C.")
    print("  " + "-" * 60 + "\n")
    if os.environ.get("TARAYICI_AC") == "1":
        _hazir_olunca_tarayici_ac(port)
    try:
        from waitress import serve
        serve(app, host=host, port=port, threads=8)
    except ImportError:
        app.run(host=host, port=port, debug=False)
