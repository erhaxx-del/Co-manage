# -*- coding: utf-8 -*-
"""Duman testleri: giriş, yetki, CRUD ve tüm sayfaların açılması."""
import os, re, sys, tempfile
import pytest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from app import uygulama_olustur

SIFRE = "Deneme2026"


@pytest.fixture()
def app():
    tmp = tempfile.mkdtemp()
    a = uygulama_olustur({"TESTING": True, "VERITABANI": os.path.join(tmp, "t.db"),
                          "SIFRE_DOSYASI": os.path.join(tmp, "s.txt"), "SABIT_SIFRE": SIFRE,
                          "HAZIR_KULLANICILAR": True})
    return a


def csrf(c, url="/giris"):
    html = c.get(url).get_data(as_text=True)
    m = re.search(r'name="_csrf" value="([^"]+)"', html)
    return m.group(1) if m else ""


def giris(c, ka, yeni="YeniSifre1"):
    t = csrf(c)
    r = c.post("/giris", data={"kullanici_adi": ka, "sifre": SIFRE, "_csrf": t})
    assert r.status_code == 302
    r = c.get("/")
    assert r.status_code == 302 and "/sifre" in r.headers["Location"]  # ilk girişte şifre değiştirme zorunlu
    t = csrf(c, "/sifre")
    r = c.post("/sifre", data={"eski": SIFRE, "yeni": yeni, "yeni2": yeni, "_csrf": t})
    assert r.status_code == 302
    return c


def test_giris_gerekli(app):
    c = app.test_client()
    assert c.get("/").status_code == 302
    assert c.get("/projeler").status_code == 302


def test_hatali_sifre_ve_kilit(app):
    c = app.test_client()
    for _ in range(5):
        t = csrf(c)
        r = c.post("/giris", data={"kullanici_adi": "salih", "sifre": "yanlis", "_csrf": t})
    t = csrf(c)
    r = c.post("/giris", data={"kullanici_adi": "salih", "sifre": SIFRE, "_csrf": t}, follow_redirects=True)
    assert "kilitlendi" in r.get_data(as_text=True)


def test_csrf_olmadan_post_reddedilir(app):
    c = app.test_client()
    r = c.post("/giris", data={"kullanici_adi": "erhan", "sifre": SIFRE})
    assert r.status_code == 400


def test_tum_sayfalar_admin(app):
    c = giris(app.test_client(), "erhan")
    for url in ["/", "/projeler", "/projeler/NW14-1", "/projeler/NW14-1/duzenle", "/projeler/yeni",
                "/projeler/yeni?ust=NW14-1", "/projeler/yeni?aksaklik=1", "/ziyaretler", "/ziyaretler?gorunum=liste",
                "/ziyaretler?hafta=1", "/ziyaretler/yeni", "/ziyaretler/1/duzenle", "/bayiler", "/bayiler/1",
                "/bayiler?sirala=risk", "/aksakliklar", "/aksakliklar?durum=", "/mintika", "/pazar-payi", "/kilavuz",
                "/yonetim/kullanicilar", "/yonetim/islem-gecmisi", "/api/kanopiler?sm=SALİH KOSİF",
                "/api/bayiler?sm=SALİH KOSİF", "/api/sonraki-id?sm=SALİH KOSİF", "/api/sonraki-id?ust=NW14-1"]:
        r = c.get(url)
        assert r.status_code == 200, (url, r.status_code)
    assert c.get("/api/sonraki-id?sm=SALİH KOSİF").json["id"] == "NW14-4"
    assert c.get("/api/sonraki-id?ust=NW14-1").json["id"] == "NW14-14"
    r = c.get("/disa-aktar"); assert r.status_code == 200 and r.data[:2] == b"PK"
    r = c.get("/yonetim/yedek"); assert r.status_code == 200 and r.data[:15] == b"SQLite format 3"


def test_saha_yetkileri(app):
    c = giris(app.test_client(), "salih")
    assert c.get("/").status_code == 200
    assert c.get("/projeler/NW11-1").status_code == 200          # başkasınınkini görebilir
    assert c.get("/projeler/NW11-1/duzenle").status_code == 403  # ama düzenleyemez
    assert c.get("/projeler/NW14-1/duzenle").status_code == 200  # kendisininkini düzenler
    assert c.get("/yonetim/kullanicilar").status_code == 403
    t = csrf(c, "/projeler/NW14-1")  # geçerli jeton, ama başkasının projesi
    assert c.post("/projeler/NW11-1/gelisme", data={"metin": "x", "_csrf": t}).status_code == 403
    # kendi projesine gelişme ekler
    t = csrf(c, "/projeler/NW14-2")
    r = c.post("/projeler/NW14-2/gelisme", data={"metin": "Test gelişmesi", "asama": "İlerleme", "ilerleme": "70",
                                                  "ilerleme_yansit": "1", "_csrf": t})
    assert r.status_code == 302
    assert "Test gelişmesi" in c.get("/projeler/NW14-2").get_data(as_text=True)


def test_proje_olustur_ve_alt_adim(app):
    c = giris(app.test_client(), "salih")
    t = csrf(c, "/projeler/yeni")
    veri = {"ad": "Test projesi", "departman": "Satış", "oncelik": "Yüksek", "durum": "Yapılacak",
            "termin": "2026-12-31", "ilerleme": "0", "_csrf": t, "sm_adi": "OZAN SELBUZ"}  # başkası adına deneme
    r = c.post("/projeler/yeni", data=veri)
    assert r.status_code == 302 and "NW14-4" in r.headers["Location"]  # saha müdürü kendi koduna zorlanır
    t = csrf(c, "/projeler/yeni?ust=NW14-4")
    r = c.post("/projeler/yeni", data=dict(veri, ust_id="NW14-4", ad="Alt adım", _csrf=t))
    assert r.status_code == 302 and "NW14-41" in r.headers["Location"]
    t = csrf(c, "/projeler/NW14-4")
    r = c.post("/projeler/NW14-4/sil", data={"_csrf": t}, follow_redirects=True)
    assert "silinemez" in r.get_data(as_text=True)


def test_ziyaret_baskasinin_bayisi_reddedilir(app):
    c = giris(app.test_client(), "salih")
    import sqlite3
    db = sqlite3.connect(app.config["VERITABANI"])
    baska = db.execute("SELECT id FROM bayiler WHERE sm_adi='OZAN SELBUZ' LIMIT 1").fetchone()[0]
    kendi = db.execute("SELECT id FROM bayiler WHERE sm_adi='SALİH KOSİF' LIMIT 1").fetchone()[0]
    t = csrf(c, "/ziyaretler/yeni")
    r = c.post("/ziyaretler/yeni", data={"tarih": "2026-09-25", "bayi_id": baska, "amac": "Rutin Ziyaret", "_csrf": t},
               follow_redirects=True)
    assert "mıntıkasında değil" in r.get_data(as_text=True)
    t = csrf(c, "/ziyaretler/yeni")
    r = c.post("/ziyaretler/yeni", data={"tarih": "2026-09-25", "bayi_id": kendi, "amac": "Rutin Ziyaret", "_csrf": t})
    assert r.status_code == 302


def test_aksakliktan_proje(app):
    c = giris(app.test_client(), "erhan")
    t = csrf(c, "/projeler/yeni?aksaklik=1")
    html = c.get("/projeler/yeni?aksaklik=1").get_data(as_text=True)
    assert 'name="aksaklik_id"' in html


def test_kayit_onay_akisi(app):
    c = app.test_client()
    assert "Kayıt olun" in c.get("/giris").get_data(as_text=True)
    t = csrf(c, "/kayit")
    # hatalı: şifreler farklı, geçersiz kullanıcı adı
    r = c.post("/kayit", data={"_csrf": t, "mintika": "NW15", "ad_soyad": "Deneme Kişi", "kullanici_adi": "Ş!",
                               "sifre": "Abcdefg1", "sifre2": "Abcdefg2"})
    h = r.get_data(as_text=True)
    assert "Kullanıcı adı 3–20" in h and "aynı değil" in h
    # alınmış kullanıcı adı
    r = c.post("/kayit", data={"_csrf": t, "mintika": "NW15", "ad_soyad": "Deneme Kişi", "kullanici_adi": "bedir",
                               "sifre": "Abcdefg1", "sifre2": "Abcdefg1"})
    assert "alınmış" in r.get_data(as_text=True)
    # başarılı kayıt
    r = c.post("/kayit", data={"_csrf": t, "mintika": "NW15", "ad_soyad": "Yeni Saha", "kullanici_adi": "yeni.saha",
                               "sifre": "Abcdefg1", "sifre2": "Abcdefg1"})
    assert "Kaydınız alındı" in r.get_data(as_text=True)
    # onaysız giriş engellenir
    t = csrf(c)
    r = c.post("/giris", data={"kullanici_adi": "yeni.saha", "sifre": "Abcdefg1", "_csrf": t}, follow_redirects=True)
    assert "onaylanmadı" in r.get_data(as_text=True)
    # bölge müdürü onaylar
    a = giris(app.test_client(), "erhan")
    h = a.get("/yonetim/kullanicilar").get_data(as_text=True)
    assert "Onay bekleyen kayıtlar" in h and "yeni.saha" in h and "zaten aktif hesap var" in h
    import sqlite3
    kid = sqlite3.connect(app.config["VERITABANI"]).execute(
        "SELECT id FROM kullanicilar WHERE kullanici_adi='yeni.saha'").fetchone()[0]
    t = csrf(a, "/yonetim/kullanicilar")
    a.post("/yonetim/kullanicilar", data={"_csrf": t, "id": kid, "islem": "onayla", "rol": "saha", "mintika": "NW15"})
    # onay sonrası giriş + yetki: yalnızca NW15
    c2 = app.test_client()
    t = csrf(c2)
    r = c2.post("/giris", data={"kullanici_adi": "yeni.saha", "sifre": "Abcdefg1", "_csrf": t})
    assert r.status_code == 302
    assert c2.get("/").status_code == 200                       # şifre değiştirme zorunlu değil (kendi seçti)
    assert c2.get("/projeler/NW15-1/duzenle").status_code == 200
    assert c2.get("/projeler/NW14-1/duzenle").status_code == 403


def test_bolge_muduru_kaydi(app):
    c = app.test_client()
    h = c.get("/kayit").get_data(as_text=True)
    assert "NW1 · Erhan Başaçar" in h and "Ziyaretçi" in h
    t = csrf(c, "/kayit")
    base = {"_csrf": t, "mintika": "NW1", "kullanici_adi": "erhan.basacar", "sifre": "Marmara2026", "sifre2": "Marmara2026"}
    r = c.post("/kayit", data={**base, "kurulum_kodu": "yanlis"})
    assert "Kurulum kodu hatalı" in r.get_data(as_text=True)
    r = c.post("/kayit", data={**base, "kurulum_kodu": SIFRE})
    assert r.status_code == 302
    t = csrf(c)
    r = c.post("/giris", data={"kullanici_adi": "erhan.basacar", "sifre": "Marmara2026", "_csrf": t})
    assert r.status_code == 302 and c.get("/").status_code == 200       # onay/şifre değiştirme gerekmez
    assert c.get("/yonetim/kullanicilar").status_code == 200
    # ikinci kez sahiplenilemez
    c3 = app.test_client()
    assert "hesabı oluşturuldu" in c3.get("/kayit").get_data(as_text=True)
    t = csrf(c3, "/kayit")
    r = c3.post("/kayit", data={**base, "_csrf": t, "kullanici_adi": "sahte", "kurulum_kodu": SIFRE})
    assert "zaten oluşturulmuş" in r.get_data(as_text=True)


def test_ziyaretci_salt_okunur(app):
    import sqlite3
    c = app.test_client()
    t = csrf(c, "/kayit")
    r = c.post("/kayit", data={"_csrf": t, "mintika": "ZIYARETCI", "ad_soyad": "Misafir Kişi", "kullanici_adi": "misafir",
                               "sifre": "Abcdefg1", "sifre2": "Abcdefg1"})
    assert "Ziyaretçi" in r.get_data(as_text=True)
    a = giris(app.test_client(), "erhan")
    kid = sqlite3.connect(app.config["VERITABANI"]).execute(
        "SELECT id FROM kullanicilar WHERE kullanici_adi='misafir'").fetchone()[0]
    t = csrf(a, "/yonetim/kullanicilar")
    a.post("/yonetim/kullanicilar", data={"_csrf": t, "id": kid, "islem": "onayla", "rol": "ziyaretci", "mintika": "NW11"})
    z = app.test_client()
    t = csrf(z)
    assert z.post("/giris", data={"kullanici_adi": "misafir", "sifre": "Abcdefg1", "_csrf": t}).status_code == 302
    for url in ["/", "/projeler", "/projeler/NW15-1", "/ziyaretler", "/bayiler", "/aksakliklar", "/mintika", "/pazar-payi"]:
        assert z.get(url).status_code == 200, url
    h = z.get("/projeler").get_data(as_text=True)
    assert "+ Yeni proje" not in h and "salt okunur" in h
    assert z.get("/projeler/yeni").status_code == 403
    assert z.get("/projeler/NW15-1/duzenle").status_code == 403
    assert z.get("/ziyaretler/yeni").status_code == 403
    assert z.get("/disa-aktar").status_code == 403
    assert z.get("/yonetim/kullanicilar").status_code == 403
    assert z.get("/isler").status_code == 403
    t = csrf(z, "/sifre")
    assert z.post("/projeler/NW15-1/gelisme", data={"_csrf": t, "aciklama": "x"}).status_code == 403


def test_eski_veritabani_gocu(tmp_path):
    import sqlite3, db as vt
    yol = tmp_path / "eski.db"
    con = sqlite3.connect(yol)
    con.executescript(vt.SEMA.replace("'admin','saha','ziyaretci'", "'admin','saha'"))
    con.execute("INSERT INTO kullanicilar (kullanici_adi, ad_soyad, rol, sifre_hash) VALUES ('x','X','admin','h')")
    con.commit()
    vt.goc(con)
    con.execute("INSERT INTO kullanicilar (kullanici_adi, ad_soyad, rol, sifre_hash) VALUES ('z','Z','ziyaretci','h')")
    assert con.execute("SELECT COUNT(*) FROM kullanicilar").fetchone()[0] == 2


def _db(app):
    import sqlite3
    c = sqlite3.connect(app.config["VERITABANI"]); c.row_factory = sqlite3.Row
    return c


def test_bolge_mudurune_is_atama(app):
    s = giris(app.test_client(), "salih")          # NW14
    t = csrf(s, "/isler/yeni")
    r = s.post("/isler/yeni", data={"_csrf": t, "baslik": "Sakarya bayisi iskonto onayı", "aciklama": "0,40 TL/lt",
                                    "oncelik": "Yüksek", "termin": "2026-10-01", "departman": ""})
    assert r.status_code == 302
    tid = _db(app).execute("SELECT id FROM talepler").fetchone()[0]
    # başka mıntıkanın projesi/bayisi seçilemez
    bayi15 = _db(app).execute("SELECT id FROM bayiler WHERE sm_adi='BEDİR EREN KÖSEOĞLU'").fetchone()[0]
    r = s.post("/isler/yeni", data={"_csrf": t, "baslik": "Deneme iş", "oncelik": "Orta", "bayi_id": bayi15})
    assert "kendi mıntıkanızdan" in r.get_data(as_text=True)
    # diğer saha müdürü göremez, ziyaretçi erişemez
    b = giris(app.test_client(), "bedir")
    assert b.get(f"/isler/{tid}").status_code == 403
    assert "iskonto onayı" not in b.get("/isler").get_data(as_text=True)
    # bölge müdürü: panoda ve ayrı sayfada görür, rozet
    a = giris(app.test_client(), "erhan")
    assert "iskonto onayı" in a.get("/").get_data(as_text=True)
    h = a.get("/isler").get_data(as_text=True)
    assert "iskonto onayı" in h and "rozet-say" in h
    assert a.get("/isler/yeni").status_code == 403              # BM kendine iş atamaz
    t = csrf(a, f"/isler/{tid}")
    a.post(f"/isler/{tid}/durum", data={"_csrf": t, "durum": "Tamamlandı", "yanit": "Onaylandı"})
    rr = _db(app).execute("SELECT durum, yanit, tamamlanma FROM talepler WHERE id=?", (tid,)).fetchone()
    assert rr["durum"] == "Tamamlandı" and rr["yanit"] == "Onaylandı" and rr["tamamlanma"]
    # saha müdürü yanıtı görür; kapalı iş düzenlenemez
    h = s.get(f"/isler/{tid}").get_data(as_text=True)
    assert "Onaylandı" in h
    assert s.get(f"/isler/{tid}/duzenle").status_code == 302
    # saha müdürü durumu değiştiremez
    t = csrf(s, "/sifre")
    assert s.post(f"/isler/{tid}/durum", data={"_csrf": t, "durum": "Yeni"}).status_code == 403
    # Excel'de sayfa var
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(a.get("/disa-aktar").data))
    assert "BM'ye Atanan İşler" in wb.sheetnames


def test_toplu_arsiv_ve_silme(app):
    s = giris(app.test_client(), "salih")
    db = _db(app)
    altlar = [r[0] for r in db.execute("SELECT id FROM projeler WHERE ust_id='NW14-1'")]
    assert altlar
    t = csrf(s, "/projeler")
    # başkasının projesi atlanır, kendi ana projesi alt adımlarıyla arşive gider
    s.post("/projeler/toplu", data={"_csrf": t, "islem": "arsivle", "sec": ["NW14-1", "NW15-1"]})
    db = _db(app)
    assert db.execute("SELECT arsiv FROM projeler WHERE id='NW14-1'").fetchone()[0] == 1
    assert all(db.execute("SELECT arsiv FROM projeler WHERE id=?", (x,)).fetchone()[0] == 1 for x in altlar)
    assert db.execute("SELECT COALESCE(arsiv,0) FROM projeler WHERE id='NW15-1'").fetchone()[0] == 0
    h = s.get("/projeler").get_data(as_text=True)
    assert "NW14-1<" not in h.replace("</span>", "<")
    assert "NW14-1" in s.get("/projeler?gorunum=arsiv").get_data(as_text=True)
    # arşivden çıkar
    s.post("/projeler/toplu", data={"_csrf": t, "islem": "arsivden_cikar", "sec": ["NW14-1"]})
    assert _db(app).execute("SELECT arsiv FROM projeler WHERE id='NW14-1'").fetchone()[0] == 0
    # toplu sil: ana proje + alt adımlar + gelişmeler
    s.post("/projeler/toplu", data={"_csrf": t, "islem": "sil", "sec": ["NW14-1", "NW14-2"]})
    db = _db(app)
    assert not db.execute("SELECT 1 FROM projeler WHERE id IN ('NW14-1','NW14-2') OR ust_id IN ('NW14-1','NW14-2')").fetchone()
    assert not db.execute("SELECT 1 FROM gelismeler WHERE proje_id LIKE 'NW14-1%' OR proje_id LIKE 'NW14-2%'").fetchone()
    # boş seçim → uyarı
    r = s.post("/projeler/toplu", data={"_csrf": t, "islem": "sil"}, follow_redirects=True)
    assert "en az bir proje" in r.get_data(as_text=True)


def test_kredi_risk(app):
    a = giris(app.test_client(), "erhan")
    h = a.get("/kredi-risk").get_data(as_text=True)
    assert "Vadesi geçmiş" in h and "Mıntıka kırılımı" in h and "Onay akışı" in h
    assert "NW11" in h and "NW16" in h
    # süzme: tek mıntıka
    h2 = a.get("/kredi-risk?sm=TUNAHAN+DEM%C4%B0RHAN").get_data(as_text=True)
    assert "NW12" in h2
    # panoda gösterge ve özet
    h3 = a.get("/").get_data(as_text=True)
    assert "Vadesi geçmiş" in h3 and "en yüksek vadesi geçmiş" in h3
    # veri tutarlılığı: mıntıka toplamları = genel toplam
    import app as uyg, db as vt, sqlite3
    with a.application.app_context():
        d = vt.baglan()
        K = uyg.kredi_hesapla(d)
        assert abs(sum(x["vadesi_gecmis"] for x in K["sm"]) - K["ozet"]["vadesi_gecmis"]) < 1
        assert K["ozet"]["borclu"] > 0 and K["ozet"]["onay"] > 0
        # her borç satırı bir saha müdürüne bağlı
        assert not d.execute("SELECT 1 FROM borc WHERE sm_adi NOT IN"
                             " ('OZAN SELBUZ','TUNAHAN DEMİRHAN','ENGİN ÜSTÜNYER','SALİH KOSİF',"
                             "'BEDİR EREN KÖSEOĞLU','YÜKSEL AYDINOL')").fetchone()
        assert not d.execute("SELECT 1 FROM risk_onaylari WHERE sm_adi IS NULL").fetchone()
        # yaş kırılımı toplamı = vadesi geçmiş
        r = d.execute("SELECT SUM(g1_7+g8_15+g16_30+g31_45+g45_ustu), SUM(vadesi_gecmis) FROM borc").fetchone()
        assert abs(r[0] - r[1]) < 1
    # saha müdürü de görebilir, ziyaretçi göremez
    s = giris(app.test_client(), "salih")
    assert s.get("/kredi-risk").status_code == 200
    # Excel'de kredi-risk sayfaları
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(a.get("/disa-aktar").data))
    assert "Kredi-Risk (Borç)" in wb.sheetnames and "Kredi-Risk (Onaylar)" in wb.sheetnames


def test_modu(app):
    """TEST_MODU=1: giriş ekranı atlanır, kullanıcı değiştirilebilir."""
    t = uygulama_olustur({"TESTING": True, "VERITABANI": app.config["VERITABANI"],
                          "SIFRE_DOSYASI": app.config["SIFRE_DOSYASI"], "TEST_MODU": True})
    c = t.test_client()
    h = c.get("/")
    assert h.status_code == 200                               # şifre sorulmadan pano açıldı
    assert "Test modu açık" in h.get_data(as_text=True)
    assert c.get("/giris").status_code == 302                 # giriş ekranı panoya yönlenir
    assert c.get("/yonetim/kullanicilar").status_code == 200  # bölge müdürü yetkisiyle
    import sqlite3
    kid = sqlite3.connect(app.config["VERITABANI"]).execute(
        "SELECT id FROM kullanicilar WHERE kullanici_adi='salih'").fetchone()[0]
    assert c.get(f"/test-kullanici/{kid}").status_code == 302
    h = c.get("/").get_data(as_text=True)
    assert "Salih Kosif" in h
    assert c.get("/yonetim/kullanicilar").status_code == 403   # saha müdürü yetkisiyle
    # test modu kapalıyken giriş zorunlu ve kullanıcı değiştirme kapalı
    c2 = app.test_client()
    assert c2.get("/").status_code == 302
    assert c2.get(f"/test-kullanici/{kid}").status_code == 404
